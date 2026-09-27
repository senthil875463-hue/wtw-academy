import os
import sqlite3
import time
import re
import random
import secrets
from datetime import datetime, timedelta
from contextlib import contextmanager

try:
    import psycopg2
    from psycopg2.extras import RealDictCursor
except Exception:
    psycopg2 = None
    RealDictCursor = None

from flask import (
    Flask,
    request,
    redirect,
    url_for,
    session,
    render_template_string,
    flash,
    jsonify,
    send_file,
)
from werkzeug.security import generate_password_hash, check_password_hash

try:
    from openpyxl import load_workbook, Workbook
except Exception:
    load_workbook = None
    Workbook = None

try:
    import qrcode
except Exception:
    qrcode = None


# ============================================================
# APP
# ============================================================

app = Flask(__name__)

app.secret_key = os.environ.get(
    "WTW_SECRET_KEY",
    "WTW_ACADEMY_SECRET_2026_STABLE"
)

app.config["SESSION_COOKIE_NAME"] = "wtw_academy_session"
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(days=30)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
INSTANCE_DIR = os.path.join(BASE_DIR, "instance")
DB_PATH = os.path.join(INSTANCE_DIR, "wtw_academy.db")
UPLOAD_DIR = os.path.join(BASE_DIR, "static", "uploads")
DATABASE_URL = os.environ.get("DATABASE_URL", "").strip()
USE_POSTGRES = bool(DATABASE_URL)

os.makedirs(INSTANCE_DIR, exist_ok=True)
os.makedirs(UPLOAD_DIR, exist_ok=True)

if USE_POSTGRES and psycopg2 is None:
    raise RuntimeError(
        "DATABASE_URL is set, but psycopg2 is not installed. "
        "Install psycopg2-binary."
    )


class DatabaseConnection:
    """Small compatibility layer so the existing WTW code can use either
    SQLite locally or PostgreSQL on Render without rewriting every route."""

    def __init__(self):
        if USE_POSTGRES:
            self.conn = psycopg2.connect(DATABASE_URL, sslmode="require")
            self.conn.autocommit = False
        else:
            self.conn = sqlite3.connect(
                DB_PATH, timeout=20, check_same_thread=False
            )
            self.conn.row_factory = sqlite3.Row

    def _sql(self, sql):
        sql = str(sql)
        if not USE_POSTGRES:
            return sql

        stripped = sql.strip()

        # SQLite PRAGMAs are not used by PostgreSQL.
        if stripped.upper().startswith("PRAGMA"):
            if re.search(r"PRAGMA\s+table_info\s*\(", stripped, re.I):
                m = re.search(r"PRAGMA\s+table_info\s*\(\s*([A-Za-z_][A-Za-z0-9_]*)\s*\)", stripped, re.I)
                table = m.group(1) if m else ""
                return (
                    "SELECT ordinal_position AS cid, column_name AS name, "
                    "data_type AS type, 0 AS notnull, column_default AS dflt_value, "
                    "ordinal_position AS pk FROM information_schema.columns "
                    "WHERE table_schema='public' AND table_name=%s "
                    "ORDER BY ordinal_position"
                ), (table,)
            return "SELECT 1 WHERE FALSE"

        # SQLite catalog query used by table_exists().
        if "sqlite_master" in stripped.lower():
            return (
                "SELECT table_name AS name FROM information_schema.tables "
                "WHERE table_schema='public' AND table_name=%s",
                None
            )

        # PostgreSQL uses SERIAL instead of SQLite AUTOINCREMENT.
        sql = re.sub(
            r"INTEGER\s+PRIMARY\s+KEY\s+AUTOINCREMENT",
            "SERIAL PRIMARY KEY",
            sql, flags=re.I
        )
        sql = re.sub(r"\bAUTOINCREMENT\b", "", sql, flags=re.I)

        # Convert SQLite placeholders.
        sql = sql.replace("?", "%s")

        # Convert INSERT OR IGNORE to PostgreSQL syntax.
        if re.match(r"^INSERT\s+OR\s+IGNORE\s+INTO\b", sql.strip(), re.I):
            sql = re.sub(
                r"^\s*INSERT\s+OR\s+IGNORE\s+INTO",
                "INSERT INTO", sql, count=1, flags=re.I
            )
            sql = sql.rstrip().rstrip(";") + " ON CONFLICT DO NOTHING"

        return sql

    def execute(self, sql, params=()):
        if not USE_POSTGRES:
            return self.conn.execute(sql, params)

        converted = self._sql(sql)
        if isinstance(converted, tuple):
            converted_sql, converted_params = converted
            if converted_params is None:
                converted_params = params
            params = converted_params
            converted = converted_sql
        cur = self.conn.cursor(cursor_factory=RealDictCursor)
        cur.execute(converted, params)
        return cur

    def executemany(self, sql, rows):
        if not USE_POSTGRES:
            return self.conn.executemany(sql, rows)
        converted = self._sql(sql)
        if isinstance(converted, tuple):
            converted = converted[0]
        cur = self.conn.cursor(cursor_factory=RealDictCursor)
        cur.executemany(converted, rows)
        return cur

    def commit(self):
        return self.conn.commit()

    def rollback(self):
        return self.conn.rollback()

    def close(self):
        return self.conn.close()


# ============================================================
# DATABASE
# ============================================================

@contextmanager
def db_conn():
    db = DatabaseConnection()
    try:
        if not USE_POSTGRES:
            db.conn.execute("PRAGMA busy_timeout=20000")
            db.conn.execute("PRAGMA foreign_keys=ON")
        yield db
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def execute_write(sql, params=(), retries=8):
    last_error = None
    for attempt in range(retries):
        try:
            with db_conn() as conn:
                cur = conn.execute(sql, params)
                if USE_POSTGRES:
                    return None
                return cur.lastrowid
        except Exception as e:
            last_error = e
            if USE_POSTGRES or "locked" not in str(e).lower():
                raise
            time.sleep(0.25 * (attempt + 1))
    raise last_error


def execute_many(sql, rows, retries=8):
    last_error = None
    for attempt in range(retries):
        try:
            with db_conn() as conn:
                conn.executemany(sql, rows)
                return
        except Exception as e:
            last_error = e
            if USE_POSTGRES or "locked" not in str(e).lower():
                raise
            time.sleep(0.25 * (attempt + 1))
    raise last_error


def query_one(sql, params=()):
    with db_conn() as conn:
        return conn.execute(sql, params).fetchone()


def query_all(sql, params=()):
    with db_conn() as conn:
        return conn.execute(sql, params).fetchall()


def column_exists(conn, table, column):
    rows = conn.execute(
        "PRAGMA table_info(" + table + ")"
    ).fetchall()

    return any(row["name"] == column for row in rows)


def add_column_if_missing(conn, table, column, definition):
    if not column_exists(conn, table, column):
        conn.execute(
            "ALTER TABLE " + table +
            " ADD COLUMN " + column + " " + definition
        )


def table_exists(conn, table):
    row = conn.execute(
        "SELECT name FROM sqlite_master "
        "WHERE type='table' AND name=?",
        (table,)
    ).fetchone()

    return row is not None


def migrate_database():
    with db_conn() as conn:

        # ----------------------------------------------------
        # Admin
        # ----------------------------------------------------
        conn.execute("""
            CREATE TABLE IF NOT EXISTS admins (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE,
                password TEXT,
                password_hash TEXT,
                full_name TEXT DEFAULT '',
                status TEXT DEFAULT 'Active',
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)

        # Old admin_users compatibility
        if table_exists(conn, "admin_users"):
            add_column_if_missing(
                conn, "admin_users", "password", "TEXT DEFAULT ''"
            )
            add_column_if_missing(
                conn, "admin_users", "password_hash", "TEXT DEFAULT ''"
            )

        add_column_if_missing(
            conn, "admins", "password", "TEXT DEFAULT ''"
        )
        add_column_if_missing(
            conn, "admins", "password_hash", "TEXT DEFAULT ''"
        )
        add_column_if_missing(
            conn, "admins", "full_name", "TEXT DEFAULT ''"
        )
        add_column_if_missing(
            conn, "admins", "status", "TEXT DEFAULT 'Active'"
        )

        # ----------------------------------------------------
        # Students
        # ----------------------------------------------------
        conn.execute("""
            CREATE TABLE IF NOT EXISTS students (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id TEXT UNIQUE,
                student_name TEXT DEFAULT '',
                name TEXT DEFAULT '',
                email TEXT DEFAULT '',
                phone TEXT DEFAULT '',
                password TEXT DEFAULT '',
                password_hash TEXT DEFAULT '',
                status TEXT DEFAULT 'Active',
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)

        add_column_if_missing(
            conn, "students", "student_name", "TEXT DEFAULT ''"
        )
        add_column_if_missing(
            conn, "students", "name", "TEXT DEFAULT ''"
        )
        add_column_if_missing(
            conn, "students", "email", "TEXT DEFAULT ''"
        )
        add_column_if_missing(
            conn, "students", "phone", "TEXT DEFAULT ''"
        )
        add_column_if_missing(
            conn, "students", "password", "TEXT DEFAULT ''"
        )
        add_column_if_missing(
            conn, "students", "password_hash", "TEXT DEFAULT ''"
        )
        add_column_if_missing(
            conn, "students", "status", "TEXT DEFAULT 'Active'"
        )

        conn.execute("""
            UPDATE students
            SET student_name=name
            WHERE COALESCE(student_name,'')=''
            AND COALESCE(name,'')<>''
        """)

        conn.execute("""
            UPDATE students
            SET name=student_name
            WHERE COALESCE(name,'')=''
            AND COALESCE(student_name,'')<>''
        """)

        conn.execute("""
            UPDATE students
            SET password=password_hash
            WHERE COALESCE(password,'')=''
            AND COALESCE(password_hash,'')<>''
        """)

        # ----------------------------------------------------
        # Subjects
        # ----------------------------------------------------
        conn.execute("""
            CREATE TABLE IF NOT EXISTS subjects (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                subject_code TEXT UNIQUE,
                subject_name TEXT,
                code TEXT DEFAULT '',
                name TEXT DEFAULT '',
                status TEXT DEFAULT 'Active',
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)

        add_column_if_missing(
            conn, "subjects", "subject_code", "TEXT DEFAULT ''"
        )
        add_column_if_missing(
            conn, "subjects", "subject_name", "TEXT DEFAULT ''"
        )
        add_column_if_missing(
            conn, "subjects", "code", "TEXT DEFAULT ''"
        )
        add_column_if_missing(
            conn, "subjects", "name", "TEXT DEFAULT ''"
        )
        add_column_if_missing(
            conn, "subjects", "status", "TEXT DEFAULT 'Active'"
        )

        conn.execute("""
            UPDATE subjects
            SET subject_code=code
            WHERE COALESCE(subject_code,'')=''
            AND COALESCE(code,'')<>''
        """)

        conn.execute("""
            UPDATE subjects
            SET subject_name=name
            WHERE COALESCE(subject_name,'')=''
            AND COALESCE(name,'')<>''
        """)

        # ----------------------------------------------------
        # Questions
        # ----------------------------------------------------
        conn.execute("""
            CREATE TABLE IF NOT EXISTS questions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                question_id TEXT,
                subject_id INTEGER,
                question_text TEXT DEFAULT '',
                option_a TEXT DEFAULT '',
                option_b TEXT DEFAULT '',
                option_c TEXT DEFAULT '',
                option_d TEXT DEFAULT '',
                correct_answer TEXT DEFAULT 'A',
                marks REAL DEFAULT 1,
                negative_mark REAL DEFAULT 0,
                difficulty TEXT DEFAULT 'Medium',
                topic TEXT DEFAULT '',
                question_type TEXT DEFAULT 'MCQ',
                status TEXT DEFAULT 'Active',
                explanation TEXT DEFAULT '',
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)

        for col, definition in [
            ("question_id", "TEXT DEFAULT ''"),
            ("subject_id", "INTEGER"),
            ("question_text", "TEXT DEFAULT ''"),
            ("option_a", "TEXT DEFAULT ''"),
            ("option_b", "TEXT DEFAULT ''"),
            ("option_c", "TEXT DEFAULT ''"),
            ("option_d", "TEXT DEFAULT ''"),
            ("correct_answer", "TEXT DEFAULT 'A'"),
            ("marks", "REAL DEFAULT 1"),
            ("negative_mark", "REAL DEFAULT 0"),
            ("difficulty", "TEXT DEFAULT 'Medium'"),
            ("topic", "TEXT DEFAULT ''"),
            ("question_type", "TEXT DEFAULT 'MCQ'"),
            ("status", "TEXT DEFAULT 'Active'"),
            ("explanation", "TEXT DEFAULT ''"),
        ]:
            add_column_if_missing(
                conn, "questions", col, definition
            )

        # ----------------------------------------------------
        # Batches
        # ----------------------------------------------------
        conn.execute("""
            CREATE TABLE IF NOT EXISTS batches (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                batch_name TEXT,
                name TEXT DEFAULT '',
                description TEXT DEFAULT '',
                status TEXT DEFAULT 'Active',
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)

        add_column_if_missing(
            conn, "batches", "batch_name", "TEXT DEFAULT ''"
        )
        add_column_if_missing(
            conn, "batches", "batch_code", "TEXT DEFAULT ''"
        )
        add_column_if_missing(
            conn, "batches", "name", "TEXT DEFAULT ''"
        )
        add_column_if_missing(
            conn, "batches", "description", "TEXT DEFAULT ''"
        )
        add_column_if_missing(
            conn, "batches", "status", "TEXT DEFAULT 'Active'"
        )

        conn.execute("""
            UPDATE batches
            SET batch_name=name
            WHERE COALESCE(batch_name,'')=''
            AND COALESCE(name,'')<>''
        """)

        conn.execute("""
            UPDATE batches
            SET batch_code = 'BATCH-' || id
            WHERE COALESCE(batch_code,'')=''
        """)

        # ----------------------------------------------------
        # Batch students
        # ----------------------------------------------------
        conn.execute("""
            CREATE TABLE IF NOT EXISTS batch_students (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                batch_id INTEGER NOT NULL,
                student_id INTEGER NOT NULL,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(batch_id, student_id)
            )
        """)

        # ----------------------------------------------------
        # Exams
        # ----------------------------------------------------
        conn.execute("""
            CREATE TABLE IF NOT EXISTS exams (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                exam_id TEXT,
                exam_name TEXT,
                name TEXT DEFAULT '',
                subject_id INTEGER,
                number_of_questions INTEGER DEFAULT 0,
                marks_per_question REAL DEFAULT 1,
                duration_minutes INTEGER DEFAULT 30,
                fee REAL DEFAULT 0,
                attempts_allowed INTEGER DEFAULT 1,
                max_attempts INTEGER DEFAULT 1,
                status TEXT DEFAULT 'Active',
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)

        for col, definition in [
            ("exam_id", "TEXT DEFAULT ''"),
            ("exam_name", "TEXT DEFAULT ''"),
            ("name", "TEXT DEFAULT ''"),
            ("subject_id", "INTEGER"),
            ("number_of_questions", "INTEGER DEFAULT 0"),
            ("marks_per_question", "REAL DEFAULT 1"),
            ("duration_minutes", "INTEGER DEFAULT 30"),
            ("fee", "REAL DEFAULT 0"),
            ("attempts_allowed", "INTEGER DEFAULT 1"),
            ("max_attempts", "INTEGER DEFAULT 1"),
            ("status", "TEXT DEFAULT 'Active'"),
        ]:
            add_column_if_missing(
                conn, "exams", col, definition
            )

        conn.execute("""
            UPDATE exams
            SET exam_name=name
            WHERE COALESCE(exam_name,'')=''
            AND COALESCE(name,'')<>''
        """)

        conn.execute("""
            UPDATE exams
            SET max_attempts=attempts_allowed
            WHERE attempts_allowed IS NOT NULL
        """)

        # ----------------------------------------------------
        # Direct exam assignment
        # ----------------------------------------------------
        conn.execute("""
            CREATE TABLE IF NOT EXISTS exam_students (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                exam_id INTEGER NOT NULL,
                student_id INTEGER NOT NULL,
                payment_required INTEGER DEFAULT 0,
                payment_status TEXT DEFAULT 'Not Required',
                payment_reference TEXT DEFAULT '',
                paid_at TEXT DEFAULT '',
                assigned_by INTEGER DEFAULT 0,
                assigned_at TEXT DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(exam_id, student_id)
            )
        """)

        for col, definition in [
            ("payment_required", "INTEGER DEFAULT 0"),
            ("payment_status", "TEXT DEFAULT 'Not Required'"),
            ("payment_reference", "TEXT DEFAULT ''"),
            ("paid_at", "TEXT DEFAULT ''"),
            ("assigned_by", "INTEGER DEFAULT 0"),
            ("assigned_at", "TEXT DEFAULT ''"),
        ]:
            add_column_if_missing(
                conn, "exam_students", col, definition
            )

        # ----------------------------------------------------
        # Batch exam assignment
        # ----------------------------------------------------
        conn.execute("""
            CREATE TABLE IF NOT EXISTS exam_batches (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                exam_id INTEGER NOT NULL,
                batch_id INTEGER NOT NULL,
                payment_required INTEGER DEFAULT 0,
                assigned_by INTEGER DEFAULT 0,
                assigned_at TEXT DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(exam_id, batch_id)
            )
        """)

        # ----------------------------------------------------
        # Attempts
        # ----------------------------------------------------
        conn.execute("""
            CREATE TABLE IF NOT EXISTS exam_attempts (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                exam_id INTEGER NOT NULL,
                student_id INTEGER NOT NULL,
                attempt_no INTEGER DEFAULT 1,
                status TEXT DEFAULT 'In Progress',
                question_order TEXT DEFAULT '',
                current_index INTEGER DEFAULT 0,
                current_question INTEGER DEFAULT 0,
                last_question INTEGER DEFAULT 0,
                started_at TEXT DEFAULT CURRENT_TIMESTAMP,
                submitted_at TEXT DEFAULT '',
                elapsed_seconds INTEGER DEFAULT 0,
                score REAL DEFAULT 0,
                total_marks REAL DEFAULT 0,
                UNIQUE(exam_id, student_id, attempt_no)
            )
        """)

        for col, definition in [
            ("attempt_no", "INTEGER DEFAULT 1"),
            ("status", "TEXT DEFAULT 'In Progress'"),
            ("question_order", "TEXT DEFAULT ''"),
            ("current_index", "INTEGER DEFAULT 0"),
            ("current_question", "INTEGER DEFAULT 0"),
            ("last_question", "INTEGER DEFAULT 0"),
            ("started_at", "TEXT DEFAULT ''"),
            ("submitted_at", "TEXT DEFAULT ''"),
            ("elapsed_seconds", "INTEGER DEFAULT 0"),
            ("score", "REAL DEFAULT 0"),
            ("total_marks", "REAL DEFAULT 0"),
        ]:
            add_column_if_missing(
                conn, "exam_attempts", col, definition
            )

        # ----------------------------------------------------
        # Answers
        # ----------------------------------------------------
        conn.execute("""
            CREATE TABLE IF NOT EXISTS exam_answers (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                attempt_id INTEGER NOT NULL,
                question_id INTEGER NOT NULL,
                selected_answer TEXT DEFAULT '',
                is_correct INTEGER DEFAULT 0,
                marks_awarded REAL DEFAULT 0,
                marked_for_review INTEGER DEFAULT 0,
                answered_at TEXT DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(attempt_id, question_id)
            )
        """)

        for col, definition in [
            ("selected_answer", "TEXT DEFAULT ''"),
            ("is_correct", "INTEGER DEFAULT 0"),
            ("marks_awarded", "REAL DEFAULT 0"),
            ("marked_for_review", "INTEGER DEFAULT 0"),
            ("answered_at", "TEXT DEFAULT ''"),
        ]:
            add_column_if_missing(
                conn, "exam_answers", col, definition
            )

 
        # ----------------------------------------------------
        # Payments
        # ----------------------------------------------------
        conn.execute("""
            CREATE TABLE IF NOT EXISTS payments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                exam_id INTEGER NOT NULL,
                student_id INTEGER NOT NULL,
                attempt_id INTEGER DEFAULT 0,
                amount REAL DEFAULT 0,
                upi_reference TEXT DEFAULT '',
                transaction_id TEXT DEFAULT '',
                status TEXT DEFAULT 'Pending',
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                verified_at TEXT DEFAULT ''
            )
        """)

        # ----------------------------------------------------
        # Payments migration
        # ----------------------------------------------------

        add_column_if_missing(
            conn,
            "payments",
            "attempt_id",
            "INTEGER DEFAULT 0"
        )

        add_column_if_missing(
            conn,
            "payments",
            "upi_reference",
            "TEXT DEFAULT ''"
        )

        add_column_if_missing(
            conn,
            "payments",
            "transaction_id",
            "TEXT DEFAULT ''"
        )

        add_column_if_missing(
            conn,
            "payments",
            "status",
            "TEXT DEFAULT 'Pending'"
        )

        add_column_if_missing(
            conn,
            "payments",
            "created_at",
            "TEXT DEFAULT ''"
        )

        add_column_if_missing(
            conn,
            "payments",
            "verified_at",
            "TEXT DEFAULT ''"
        )

        # ----------------------------------------------------
        # Indexes
        # ----------------------------------------------------
        conn.execute("""
            CREATE INDEX IF NOT EXISTS idx_questions_subject
            ON questions(subject_id)
        """)

        conn.execute("""
            CREATE INDEX IF NOT EXISTS idx_exam_students_exam
            ON exam_students(exam_id)
        """)

        conn.execute("""
            CREATE INDEX IF NOT EXISTS idx_exam_students_student
            ON exam_students(student_id)
        """)

        conn.execute("""
            CREATE INDEX IF NOT EXISTS idx_attempts_student
            ON exam_attempts(student_id)
        """)

        # ----------------------------------------------------
        # Default admin
        # ----------------------------------------------------
        admin = conn.execute(
            "SELECT id FROM admins WHERE username=?",
            ("admin",)
        ).fetchone()

        if admin is None:
            conn.execute("""
                INSERT INTO admins
                (username, password, password_hash, full_name, status)
                VALUES (?, ?, ?, ?, ?)
            """, (
                "admin",
                generate_password_hash("admin123"),
                generate_password_hash("admin123"),
                "Administrator",
                "Active"
            ))

        # ----------------------------------------------------
        # WAL
        # ----------------------------------------------------
        try:
            conn.execute("PRAGMA journal_mode=WAL")
        except sqlite3.OperationalError:
            pass

        try:
            conn.execute("PRAGMA synchronous=NORMAL")
        except sqlite3.OperationalError:
            pass


# ============================================================
# OPTIONAL SQLITE -> POSTGRES DATA MIGRATION
# ============================================================

def migrate_local_sqlite_to_postgres():
    """Copy existing local SQLite data into PostgreSQL once.

    Enable with WTW_MIGRATE_SQLITE=1. This is intentionally opt-in so a
    normal Render restart never overwrites or duplicates production data.
    """
    if not USE_POSTGRES or os.environ.get("WTW_MIGRATE_SQLITE") != "1":
        return
    if not os.path.exists(DB_PATH):
        print("SQLite migration skipped: local SQLite database not found.")
        return

    with db_conn() as pg: 
        pg.execute("""
            CREATE TABLE IF NOT EXISTS wtw_migration_log (
                id SERIAL PRIMARY KEY,
                migration_key TEXT UNIQUE NOT NULL,
                completed_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)
        done = pg.execute(
            "SELECT id FROM wtw_migration_log WHERE migration_key=%s",
            ("sqlite_v1",)
        ).fetchone()
        if done:
            print("SQLite migration already completed.")
            return

        src = sqlite3.connect(DB_PATH)
        src.row_factory = sqlite3.Row
        try:
            tables = [
                r[0] for r in src.execute(
                    "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'"
                ).fetchall()
            ]
            # Only copy tables that the current PostgreSQL schema owns.
            for table in tables:
                cols = [
                    r["name"] for r in src.execute(f"PRAGMA table_info({table})").fetchall()
                ]
                if not cols:
                    continue
                pg_cols = pg.execute(
                    "SELECT column_name FROM information_schema.columns "
                    "WHERE table_schema='public' AND table_name=%s ORDER BY ordinal_position",
                    (table,)
                ).fetchall()
                pg_col_names = [r["column_name"] for r in pg_cols]
                if not pg_col_names:
                    continue
                use_cols = [c for c in cols if c in pg_col_names]
                if not use_cols:
                    continue
                rows = src.execute(
                    "SELECT " + ",".join('"'+c.replace('"','""')+'"' for c in use_cols) +
                    " FROM " + '"'+table.replace('"','""')+'"'
                ).fetchall()
                if not rows:
                    continue
                placeholders = ",".join(["%s"] * len(use_cols))
                quoted_cols = ",".join('"'+c.replace('"','""')+'"' for c in use_cols)
                sql = f"INSERT INTO \"{table}\" ({quoted_cols}) VALUES ({placeholders}) ON CONFLICT DO NOTHING"
                pg.executemany(sql, [tuple(row[c] for c in use_cols) for row in rows])

            pg.execute(
                "INSERT INTO wtw_migration_log (migration_key) VALUES (%s) ON CONFLICT DO NOTHING",
                ("sqlite_v1",)
            )
            print("SQLite data migration completed successfully.")
        finally:
            src.close()


# ============================================================
# HELPERS
# ============================================================

def now_text():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def make_password(raw):
    return generate_password_hash(str(raw))


def verify_password(stored, provided):
    if not stored:
        return False

    try:
        if stored.startswith(("pbkdf2:", "scrypt:")):
            return check_password_hash(stored, provided)
    except Exception:
        pass

    return secrets.compare_digest(
        str(stored),
        str(provided)
    )


def safe_float(value, default=0):
    try:
        return float(value)
    except Exception:
        return default


def safe_int(value, default=0):
    try:
        return int(value)
    except Exception:
        return default


def normalize_answer(value):
    if value is None:
        return ""

    value = str(value).strip().upper()

    mapping = {
        "OPTION A": "A",
        "OPTION B": "B",
        "OPTION C": "C",
        "OPTION D": "D",
        "1": "A",
        "2": "B",
        "3": "C",
        "4": "D",
    }

    if value in mapping:
        return mapping[value]

    if "," in value:
        parts = [
            normalize_answer(x)
            for x in value.split(",")
        ]
        parts = [x for x in parts if x]
        return "|".join(sorted(set(parts)))

    if "|" in value:
        parts = [
            normalize_answer(x)
            for x in value.split("|")
        ]
        parts = [x for x in parts if x]
        return "|".join(sorted(set(parts)))

    if value in ("A", "B", "C", "D"):
        return value

    return value


def answer_is_correct(selected, correct, question_type):
    selected_set = set(
        x for x in normalize_answer(selected).split("|") if x
    )

    correct_set = set(
        x for x in normalize_answer(correct).split("|") if x
    )

    if question_type == "Multiple Correct":
        return selected_set == correct_set

    return selected_set == correct_set and len(selected_set) == 1


def admin_logged():
    return bool(session.get("admin_id"))


def student_logged():
    return bool(session.get("student_db_id"))


def admin_required():
    if not admin_logged():
        return redirect(url_for("admin_login"))
    return None


def student_required():
    if not student_logged():
        return redirect(url_for("student_login"))
    return None


# ============================================================
# TEMPLATE
# ============================================================

BASE_HTML = """
<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport"
      content="width=device-width,initial-scale=1,maximum-scale=1">
<title>{{ title }} - WTW Academy</title>

<style>
*{box-sizing:border-box}
body{
    margin:0;
    font-family:Arial,Helvetica,sans-serif;
    background:#f4f7fb;
    color:#172033;
}
a{text-decoration:none;color:inherit}
.nav{
    background:#172554;
    color:white;
    padding:12px 16px;
    display:flex;
    align-items:center;
    gap:10px;
    flex-wrap:wrap;
}
.brand{
    font-size:20px;
    font-weight:800;
    margin-right:auto;
}
.nav a{
    padding:8px 10px;
    border-radius:7px;
    font-size:14px;
}
.nav a:hover{background:rgba(255,255,255,.15)}
.container{
    width:min(1200px,94%);
    margin:22px auto;
}
.card{
    background:white;
    border-radius:14px;
    padding:18px;
    margin-bottom:18px;
    box-shadow:0 4px 18px rgba(0,0,0,.06);
}
h1{font-size:25px;margin:0 0 18px}
h2{font-size:20px;margin:0 0 14px}
h3{margin:0 0 10px}
.grid{
    display:grid;
    grid-template-columns:repeat(4,1fr);
    gap:14px;
}
.stat{
    background:white;
    padding:18px;
    border-radius:13px;
    box-shadow:0 4px 18px rgba(0,0,0,.05);
}
.stat b{
    display:block;
    font-size:27px;
    margin-top:7px;
}
label{
    display:block;
    font-weight:700;
    margin:10px 0 5px;
}
input,select,textarea{
    width:100%;
    padding:11px;
    border:1px solid #cbd5e1;
    border-radius:8px;
    font-size:15px;
    background:white;
}
textarea{min-height:100px;resize:vertical}
button,.btn{
    display:inline-block;
    border:0;
    padding:10px 15px;
    border-radius:8px;
    cursor:pointer;
    background:#2563eb;
    color:white;
    font-weight:700;
    margin:4px 2px;
}
.btn.green{background:#15803d}
.btn.red{background:#dc2626}
.btn.gray{background:#64748b}
.btn.orange{background:#ea580c}
.btn.dark{background:#334155}
.btn.small{padding:7px 10px;font-size:13px}
table{
    width:100%;
    border-collapse:collapse;
}
th,td{
    padding:10px;
    border-bottom:1px solid #e5e7eb;
    text-align:left;
    vertical-align:top;
}
th{
    background:#f8fafc;
    font-size:13px;
}
.actions{white-space:nowrap}
.badge{
    display:inline-block;
    padding:5px 9px;
    border-radius:20px;
    background:#e2e8f0;
    font-size:12px;
    font-weight:700;
}
.badge.green{background:#dcfce7;color:#166534}
.badge.red{background:#fee2e2;color:#991b1b}
.badge.orange{background:#ffedd5;color:#9a3412}
.flash{
    padding:12px 15px;
    border-radius:9px;
    background:#dbeafe;
    margin-bottom:12px;
    font-weight:700;
}
.login-wrap{
    min-height:100vh;
    display:flex;
    align-items:center;
    justify-content:center;
    padding:20px;
}
.login-card{
    width:min(430px,100%);
    background:white;
    padding:28px;
    border-radius:18px;
    box-shadow:0 10px 35px rgba(0,0,0,.1);
}
.question{
    font-size:18px;
    line-height:1.6;
    font-weight:700;
    margin-bottom:15px;
}
.option{
    border:1px solid #dbe2ea;
    border-radius:9px;
    padding:12px;
    margin:9px 0;
    cursor:pointer;
}
.option:hover{background:#f8fafc}
.option input{width:auto;margin-right:8px}
.exam-layout{
    display:grid;
    grid-template-columns:1fr 260px;
    gap:18px;
}
.navigator{
    display:grid;
    grid-template-columns:repeat(5,1fr);
    gap:6px;
}
.navigator a{
    padding:9px 4px;
    text-align:center;
    border-radius:7px;
    background:#e2e8f0;
    font-weight:700;
    font-size:13px;
}
.navigator a.current{background:#2563eb;color:white}
.navigator a.answered{background:#bbf7d0}
.review{
    background:#fef3c7!important;
}
.center{text-align:center}
.muted{color:#64748b}
.row{
    display:grid;
    grid-template-columns:repeat(2,1fr);
    gap:14px;
}
.row3{
    display:grid;
    grid-template-columns:repeat(3,1fr);
    gap:14px;
}
@media(max-width:850px){
    .grid{grid-template-columns:repeat(2,1fr)}
    .exam-layout{grid-template-columns:1fr}
}
@media(max-width:600px){
    .container{width:96%;margin:12px auto}
    .nav{padding:10px}
    .nav a{font-size:12px;padding:7px}
    .grid,.row,.row3{grid-template-columns:1fr}
    table{font-size:13px}
    th,td{padding:7px}
    .table-scroll{overflow-x:auto}
    h1{font-size:21px}
}
</style>
</head>

<body>

{% if nav %}
<div class="nav">
    <div class="brand">WTW Academy</div>

    {% if admin %}
        <a href="{{ url_for('admin_dashboard') }}">Dashboard</a>
        <a href="{{ url_for('admin_students') }}">Students</a>
        <a href="{{ url_for('admin_batches') }}">Batches</a>
        <a href="{{ url_for('admin_subjects') }}">Subjects</a>
        <a href="{{ url_for('admin_questions') }}">Questions</a>
        <a href="{{ url_for('admin_exams') }}">Exams</a>
        <a href="{{ url_for('admin_payments') }}">Payments</a>
        <a href="{{ url_for('admin_results') }}">Results</a>
        <a href="{{ url_for('admin_logout') }}">Logout</a>
    {% elif student %}
        <a href="{{ url_for('student_dashboard') }}">Dashboard</a>
        <a href="{{ url_for('student_results') }}">Results</a>
        <a href="{{ url_for('student_logout') }}">Logout</a>
    {% endif %}
</div>
{% endif %}

<div class="container">

{% with messages=get_flashed_messages() %}
{% for message in messages %}
<div class="flash">{{ message }}</div>
{% endfor %}
{% endwith %}

{{ body|safe }}

</div>
<script>
function togglePassword(id, button) {
    const input = document.getElementById(id);

    if (input.type === "password") {
        input.type = "text";
        button.innerText = "🙈";
    } else {
        input.type = "password";
        button.innerText = "👁";
    }
}
</script>

</body>
</html>
"""


def page(title, body, nav=True, admin=False, student=False, **context):
    # Keep page flags separate from template context.
    # If a route passes student=student, it is captured by the
    # function argument above, so explicitly put it back into
    # the context used to render the inner template.
    template_context = dict(context)
    template_context["admin"] = admin
    template_context["student"] = student

    rendered_body = render_template_string(
        body,
        **template_context
    )

    return render_template_string(
        BASE_HTML,
        title=title,
        body=rendered_body,
        nav=nav,
        admin=admin,
        student=student,
        **context
    )


# ============================================================
# PUBLIC
# ============================================================

@app.route("/")
def home():
    body = """
    <div class="card center">
        <h1>WTW Academy</h1>
        <p class="muted">Online Examination System</p>
        <p>
            <a class="btn" href="{{ url_for('student_login') }}">
                Student Login
            </a>
            <a class="btn dark" href="{{ url_for('admin_login') }}">
                Admin Login
            </a>
        </p>
    </div>
    """

    return page(
        "WTW Academy",
        body,
        nav=False
    )


# ============================================================
# ADMIN LOGIN
# ============================================================

@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")

        row = query_one(
            "SELECT * FROM admins WHERE username=?",
            (username,)
        )

        if row and row["status"] == "Active" and verify_password(
            row["password"] or row["password_hash"],
            password
        ):
            session.permanent = True
            session["admin_id"] = row["id"]
            session["admin_username"] = row["username"]

            return redirect(url_for("admin_dashboard"))

        flash("Invalid admin login.")

    body = """
    <div class="login-wrap">
        <div class="login-card">
            <h1>WTW Academy</h1>
            <h2>Admin Login</h2>

            <form method="post">

                <label>Username</label>
                <input name="username" required autofocus>

                <label>Password</label>

                <div style="position:relative">
                    <input type="password"
                           name="password"
                           id="adminPassword"
                           required
                           style="padding-right:45px">

                    <button type="button"
                            onclick="togglePassword('adminPassword', this)"
                            style="position:absolute;
                                   right:8px;
                                   top:50%;
                                   transform:translateY(-50%);
                                   border:0;
                                   background:none;
                                   cursor:pointer;color:#334155;font-size:20px">
                        👁
                    </button>
                </div>

                <button style="width:100%;margin-top:15px">
                    Login
                </button>

            </form>
        </div>
    </div>
    """

    return page(
        "Admin Login",
        body,
        nav=False
    )


@app.route("/admin/logout")
def admin_logout():
    session.pop("admin_id", None)
    session.pop("admin_username", None)
    return redirect(url_for("admin_login"))
# ============================================================
# ADMIN DASHBOARD
# ============================================================

@app.route("/admin")
@app.route("/admin/dashboard")
def admin_dashboard():

    guard = admin_required()
    if guard:
        return guard

    students = query_one(
        "SELECT COUNT(*) c FROM students"
    )["c"]

    subjects = query_one(
        "SELECT COUNT(*) c FROM subjects"
    )["c"]

    questions = query_one(
        "SELECT COUNT(*) c FROM questions"
    )["c"]

    exams = query_one(
        "SELECT COUNT(*) c FROM exams"
    )["c"]

    batches = query_one(
        "SELECT COUNT(*) c FROM batches"
    )["c"]

    results = query_one(
        "SELECT COUNT(*) c FROM exam_attempts "
        "WHERE status='Submitted'"
    )["c"]

    body = """
    <h1>Admin Dashboard</h1>

    <div class="grid">
        <div class="stat">Students<b>{{ students }}</b></div>
        <div class="stat">Subjects<b>{{ subjects }}</b></div>
        <div class="stat">Questions<b>{{ questions }}</b></div>
        <div class="stat">Exams<b>{{ exams }}</b></div>
        <div class="stat">Batches<b>{{ batches }}</b></div>
        <div class="stat">Results<b>{{ results }}</b></div>
    </div>

    <div class="card">
        <h2>Quick Actions</h2>
        <a class="btn" href="{{ url_for('student_add') }}">
            Add Student
        </a>
        <a class="btn" href="{{ url_for('subject_add') }}">
            Add Subject
        </a>
        <a class="btn" href="{{ url_for('question_add') }}">
            Add Question
        </a>
        <a class="btn green" href="{{ url_for('exam_add') }}">
            Create Exam
        </a>
        <a class="btn orange" href="{{ url_for('batch_add') }}">
            Create Batch
        </a>
    </div>
    """

    return page(
        "Dashboard",
        body,
        admin=True,
        students=students,
        subjects=subjects,
        questions=questions,
        exams=exams,
        batches=batches,
        results=results
    )


# ============================================================
# STUDENTS
# ============================================================

@app.route("/admin/students")
def admin_students():

    guard = admin_required()
    if guard:
        return guard

    students = query_all("""
        SELECT *
        FROM students
        ORDER BY id DESC
    """)

    body = """
    <div class="card">
        <h1>Students</h1>

        <a class="btn green"
           href="{{ url_for('student_add') }}">
            Add Student
        </a>

        <div class="table-scroll">
        <table>
        <tr>
            <th>ID</th>
            <th>Student ID</th>
            <th>Name</th>
            <th>Email</th>
            <th>Phone</th>
            <th>Status</th>
            <th>Actions</th>
        </tr>

        {% for s in students %}
        <tr>
            <td>{{ s.id }}</td>
            <td>{{ s.student_id }}</td>
            <td>{{ s.student_name or s.name }}</td>
            <td>{{ s.email }}</td>
            <td>{{ s.phone }}</td>
            <td>{{ s.status }}</td>
            <td class="actions">
                <a class="btn small"
                   href="{{ url_for('student_edit', sid=s.id) }}">
                   Edit
                </a>

                <form method="post"
                      action="{{ url_for('student_toggle', sid=s.id) }}"
                      style="display:inline">
                    <button class="btn small gray">
                        Toggle
                    </button>
                </form>

                <form method="post"
                      action="{{ url_for('student_delete', sid=s.id) }}"
                      style="display:inline"
                      onsubmit="return confirm('Delete student?')">
                    <button class="btn small red">
                        Delete
                    </button>
                </form>
            </td>
        </tr>
        {% endfor %}
        </table>
        </div>
    </div>
    """

    return page(
        "Students",
        body,
        admin=True,
        students=students
    )


@app.route("/admin/students/add", methods=["GET", "POST"])
def student_add():

    guard = admin_required()
    if guard:
        return guard

    if request.method == "POST":

        student_id = request.form.get("student_id", "").strip()
        name = request.form.get("student_name", "").strip()
        email = request.form.get("email", "").strip()
        phone = request.form.get("phone", "").strip()
        password = request.form.get("password", "")

        if not student_id or not name or not password:
            flash("Student ID, Name and Password are required.")
        else:
            exists = query_one(
                "SELECT id FROM students WHERE student_id=?",
                (student_id,)
            )

            if exists:
                flash("Student ID already exists.")
            else:
                hashed = make_password(password)

                execute_write("""
                    INSERT INTO students
                    (student_id,student_name,name,email,phone,
                     password,password_hash,status)
                    VALUES (?,?,?,?,?,?,?,?)
                """, (
                    student_id,
                    name,
                    name,
                    email,
                    phone,
                    hashed,
                    hashed,
                    "Active"
                ))

                flash("Student created successfully.")
                return redirect(url_for("admin_students"))

    body = """
    <div class="card">
        <h1>Add Student</h1>

        <form method="post">

            <label>Student ID</label>
            <input name="student_id" required>

            <label>Student Name</label>
            <input name="student_name" required>

            <label>Email</label>
            <input type="email" name="email">

            <label>Phone</label>
            <input name="phone">

            <label>Password</label>
            <input type="password" name="password" required>

            <button class="green">Create Student</button>
            <a class="btn gray"
               href="{{ url_for('admin_students') }}">
               Cancel
            </a>
        </form>
    </div>
    """

    return page(
        "Add Student",
        body,
        admin=True
    )


@app.route("/admin/students/<int:sid>/edit", methods=["GET", "POST"])
def student_edit(sid):

    guard = admin_required()
    if guard:
        return guard

    student = query_one(
        "SELECT * FROM students WHERE id=?",
        (sid,)
    )

    if not student:
        flash("Student not found.")
        return redirect(url_for("admin_students"))

    if request.method == "POST":

        student_id = request.form.get("student_id", "").strip()
        name = request.form.get("student_name", "").strip()
        email = request.form.get("email", "").strip()
        phone = request.form.get("phone", "").strip()
        password = request.form.get("password", "")

        if password:
            hashed = make_password(password)

            execute_write("""
                UPDATE students
                SET student_id=?,
                    student_name=?,
                    name=?,
                    email=?,
                    phone=?,
                    password=?,
                    password_hash=?
                WHERE id=?
            """, (
                student_id,
                name,
                name,
                email,
                phone,
                hashed,
                hashed,
                sid
            ))
        else:
            execute_write("""
                UPDATE students
                SET student_id=?,
                    student_name=?,
                    name=?,
                    email=?,
                    phone=?
                WHERE id=?
            """, (
                student_id,
                name,
                name,
                email,
                phone,
                sid
            ))

        flash("Student updated.")
        return redirect(url_for("admin_students"))

    body = """
    <div class="card">
        <h1>Edit Student</h1>

        <form method="post">

            <label>Student ID</label>
            <input name="student_id"
                   value="{{ student.student_id }}"
                   required>

            <label>Student Name</label>
            <input name="student_name"
                   value="{{ student.student_name or student.name }}"
                   required>

            <label>Email</label>
            <input name="email"
                   value="{{ student.email }}">

            <label>Phone</label>
            <input name="phone"
                   value="{{ student.phone }}">

            <label>New Password</label>

            <div style="position:relative;">
                <input type="password"
                       id="edit_student_password"
                       name="password"
                       placeholder="Leave blank to keep current password"
                       style="padding-right:55px;">

                <button
                    type="button"
                    onclick="togglePassword('edit_student_password', this)"
                    style="position:absolute; right:5px; top:5px; height:32px; width:42px; border:0; background:transparent; cursor:pointer; font-size:18px;">
                    👁
                </button>
            </div>

            <br>

            <button class="green">Save</button>
            <a class="btn gray"
               href="{{ url_for('admin_students') }}">
               Cancel
            </a>
        </form>
    </div>
    """

    return page(
        "Edit Student",
        body,
        admin=True,
        student=student
    )


@app.route("/admin/students/<int:sid>/toggle", methods=["POST"])
def student_toggle(sid):

    guard = admin_required()
    if guard:
        return guard

    row = query_one(
        "SELECT status FROM students WHERE id=?",
        (sid,)
    )

    if row:
        new_status = (
            "Inactive"
            if row["status"] == "Active"
            else "Active"
        )

        execute_write(
            "UPDATE students SET status=? WHERE id=?",
            (new_status, sid)
        )

    return redirect(url_for("admin_students"))


@app.route("/admin/students/<int:sid>/delete", methods=["POST"])
def student_delete(sid):

    guard = admin_required()
    if guard:
        return guard

    execute_write(
        "DELETE FROM batch_students WHERE student_id=?",
        (sid,)
    )

    execute_write(
        "DELETE FROM exam_students WHERE student_id=?",
        (sid,)
    )

    execute_write(
        "DELETE FROM students WHERE id=?",
        (sid,)
    )

    flash("Student deleted.")
    return redirect(url_for("admin_students"))


# ============================================================
# SUBJECTS
# ============================================================

@app.route("/admin/subjects")
def admin_subjects():

    guard = admin_required()
    if guard:
        return guard

    subjects = query_all("""
        SELECT *
        FROM subjects
        ORDER BY id DESC
    """)

    body = """
    <div class="card">
        <h1>Subjects</h1>

        <a class="btn green"
           href="{{ url_for('subject_add') }}">
            Add Subject
        </a>

        <div class="table-scroll">
        <table>
        <tr>
            <th>ID</th>
            <th>Code</th>
            <th>Subject</th>
            <th>Status</th>
            <th>Actions</th>
        </tr>

        {% for s in subjects %}
        <tr>
            <td>{{ s.id }}</td>
            <td>{{ s.subject_code }}</td>
            <td>{{ s.subject_name }}</td>
            <td>{{ s.status }}</td>
            <td>
                <a class="btn small"
                   href="{{ url_for('subject_edit', sid=s.id) }}">
                   Edit
                </a>

                <form method="post"
                      action="{{ url_for('subject_toggle', sid=s.id) }}"
                      style="display:inline">
                    <button class="btn small gray">
                        Toggle
                    </button>
                </form>

                <form method="post"
                      action="{{ url_for('subject_delete', sid=s.id) }}"
                      style="display:inline"
                      onsubmit="return confirm('Delete subject?')">
                    <button class="btn small red">
                        Delete
                    </button>
                </form>
            </td>
        </tr>
        {% endfor %}
        </table>
        </div>
    </div>
    """

    return page(
        "Subjects",
        body,
        admin=True,
        subjects=subjects
    )


@app.route("/admin/subjects/add", methods=["GET", "POST"])
def subject_add():

    guard = admin_required()
    if guard:
        return guard

    if request.method == "POST":

        code = request.form.get("subject_code", "").strip().upper()
        name = request.form.get("subject_name", "").strip()

        if not code or not name:
            flash("Code and Subject Name are required.")
        else:
            exists = query_one(
                "SELECT id FROM subjects WHERE subject_code=?",
                (code,)
            )

            if exists:
                flash("Subject code already exists.")
            else:
                execute_write("""
                    INSERT INTO subjects
                    (subject_code,subject_name,code,name,status)
                    VALUES (?,?,?,?,?)
                """, (
                    code,
                    name,
                    code,
                    name,
                    "Active"
                ))

                flash("Subject created.")
                return redirect(url_for("admin_subjects"))

    body = """
    <div class="card">
        <h1>Add Subject</h1>

        <form method="post">
            <label>Subject Code</label>
            <input name="subject_code" required>

            <label>Subject Name</label>
            <input name="subject_name" required>

            <button class="green">Create Subject</button>
            <a class="btn gray"
               href="{{ url_for('admin_subjects') }}">
               Cancel
            </a>
        </form>
    </div>
    """

    return page(
        "Add Subject",
        body,
        admin=True
    )


@app.route("/admin/subjects/<int:sid>/edit", methods=["GET", "POST"])
def subject_edit(sid):

    guard = admin_required()
    if guard:
        return guard

    subject = query_one(
        "SELECT * FROM subjects WHERE id=?",
        (sid,)
    )

    if not subject:
        flash("Subject not found.")
        return redirect(url_for("admin_subjects"))

    if request.method == "POST":

        code = request.form.get("subject_code", "").strip().upper()
        name = request.form.get("subject_name", "").strip()

        execute_write("""
            UPDATE subjects
            SET subject_code=?,
                subject_name=?,
                code=?,
                name=?
            WHERE id=?
        """, (
            code,
            name,
            code,
            name,
            sid
        ))

        flash("Subject updated.")
        return redirect(url_for("admin_subjects"))

    body = """
    <div class="card">
        <h1>Edit Subject</h1>

        <form method="post">

            <label>Subject Code</label>
            <input name="subject_code"
                   value="{{ subject.subject_code }}"
                   required>

            <label>Subject Name</label>
            <input name="subject_name"
                   value="{{ subject.subject_name }}"
                   required>

            <button class="green">Save</button>

            <a class="btn gray"
               href="{{ url_for('admin_subjects') }}">
               Cancel
            </a>
        </form>
    </div>
    """

    return page(
        "Edit Subject",
        body,
        admin=True,
        subject=subject
    )


@app.route("/admin/subjects/<int:sid>/toggle", methods=["POST"])
def subject_toggle(sid):

    guard = admin_required()
    if guard:
        return guard

    row = query_one(
        "SELECT status FROM subjects WHERE id=?",
        (sid,)
    )

    if row:
        status = (
            "Inactive"
            if row["status"] == "Active"
            else "Active"
        )

        execute_write(
            "UPDATE subjects SET status=? WHERE id=?",
            (status, sid)
        )

    return redirect(url_for("admin_subjects"))


@app.route("/admin/subjects/<int:sid>/delete", methods=["POST"])
def subject_delete(sid):

    guard = admin_required()
    if guard:
        return guard

    execute_write(
        "DELETE FROM questions WHERE subject_id=?",
        (sid,)
    )

    execute_write(
        "DELETE FROM subjects WHERE id=?",
        (sid,)
    )

    flash("Subject deleted.")
    return redirect(url_for("admin_subjects"))


# ============================================================
# QUESTIONS
# ============================================================

@app.route("/admin/questions")
def admin_questions():

    guard = admin_required()
    if guard:
        return guard

    search = request.args.get("search", "").strip()
    subject = request.args.get("subject", "")
    difficulty = request.args.get("difficulty", "")

    sql = """
        SELECT q.*, s.subject_name, s.subject_code
        FROM questions q
        LEFT JOIN subjects s
        ON q.subject_id=s.id
        WHERE 1=1
    """

    params = []

    if search:
        sql += """
        AND (
            q.question_id LIKE ?
            OR q.question_text LIKE ?
            OR q.option_a LIKE ?
            OR q.option_b LIKE ?
            OR q.option_c LIKE ?
            OR q.option_d LIKE ?
            OR q.topic LIKE ?
        )
        """

        term = "%" + search + "%"
        params.extend([
            term, term, term, term,
            term, term, term
        ])

    if subject:
        sql += " AND q.subject_id=?"
        params.append(subject)

    if difficulty:
        sql += " AND q.difficulty=?"
        params.append(difficulty)

    sql += " ORDER BY q.id DESC"

    questions = query_all(sql, params)
    subjects = query_all(
        "SELECT * FROM subjects ORDER BY subject_name"
    )

    body = """
    <div class="card">

        <h1>Question Bank</h1>

        <a class="btn green"
           href="{{ url_for('question_add') }}">
           Add Question
        </a>

        <a class="btn orange"
           href="{{ url_for('question_import') }}">
           Import Excel
        </a>

        <a class="btn gray"
           href="{{ url_for('question_template') }}">
           Excel Template
        </a>

        <form method="get" style="margin-top:15px">
            <div class="row">

                <div>
                    <label>Search</label>
                    <input name="search"
                           value="{{ search }}"
                           placeholder="Question ID / text / topic">
                </div>

                <div>
                    <label>Subject</label>
                    <select name="subject">
                        <option value="">All Subjects</option>
                        {% for s in subjects %}
                        <option value="{{ s.id }}"
                        {% if subject|string == s.id|string %}
                        selected
                        {% endif %}>
                            {{ s.subject_code }} -
                            {{ s.subject_name }}
                        </option>
                        {% endfor %}
                    </select>
                </div>

            </div>

            <label>Difficulty</label>
            <select name="difficulty">
                <option value="">All</option>
                {% for d in ["Easy","Medium","Hard"] %}
                <option value="{{ d }}"
                {% if difficulty == d %}selected{% endif %}>
                    {{ d }}
                </option>
                {% endfor %}
            </select>

            <button>Search</button>
            <a class="btn gray"
               href="{{ url_for('admin_questions') }}">
               Clear
            </a>
        </form>
    </div>

    <div class="card">

        <form method="post"
              action="{{ url_for('question_delete_selected') }}"
              onsubmit="return confirm('Delete selected questions?')">

            <button class="btn red">
                Delete Selected
            </button>

            <button type="button"
                    class="btn red"
                    onclick="selectAllQuestions()">
                Select All
            </button>
        </form>

        <form method="post"
              action="{{ url_for('question_delete_all') }}"
              onsubmit="return confirm('DELETE ALL QUESTIONS?')">
            <button class="btn red">
                Delete All Questions
            </button>
        </form>

        <div class="table-scroll">
        <table>
        <tr>
            <th>✓</th>
            <th>ID</th>
            <th>Subject</th>
            <th>Question</th>
            <th>Correct</th>
            <th>Marks</th>
            <th>Negative</th>
            <th>Difficulty</th>
            <th>Status</th>
            <th>Actions</th>
        </tr>

        {% for q in questions %}
        <tr>
            <td>
                <input class="qbox"
                       type="checkbox"
                       form="delete-selected-form"
                       name="question_ids"
                       value="{{ q.id }}">
            </td>
            <td>{{ q.question_id }}</td>
            <td>
                {{ q.subject_code }}<br>
                {{ q.subject_name }}
            </td>
            <td>{{ q.question_text }}</td>
            <td>{{ q.correct_answer }}</td>
            <td>{{ q.marks }}</td>
            <td>{{ q.negative_mark }}</td>
            <td>{{ q.difficulty }}</td>
            <td>{{ q.status }}</td>
            <td class="actions">

                <a class="btn small"
                   href="{{ url_for('question_edit', qid=q.id) }}">
                   Edit
                </a>

                <form method="post"
                      action="{{ url_for('question_toggle', qid=q.id) }}"
                      style="display:inline">
                    <button class="btn small gray">
                        Toggle
                    </button>
                </form>

                <form method="post"
                      action="{{ url_for('question_delete', qid=q.id) }}"
                      style="display:inline"
                      onsubmit="return confirm('Delete question?')">
                    <button class="btn small red">
                        Delete
                    </button>
                </form>

            </td>
        </tr>
        {% endfor %}
        </table>
        </div>
    </div>

    <form id="delete-selected-form"
          method="post"
          action="{{ url_for('question_delete_selected') }}"
          onsubmit="return confirm('Delete selected questions?')">
    </form>

    <script>
    function selectAllQuestions(){
        document.querySelectorAll('.qbox').forEach(function(x){
            x.checked=true;
        });
    }
    </script>
    """

    return page(
        "Questions",
        body,
        admin=True,
        questions=questions,
        subjects=subjects,
        search=search,
        subject=subject,
        difficulty=difficulty
    )


@app.route("/admin/questions/add", methods=["GET", "POST"])
def question_add():

    guard = admin_required()
    if guard:
        return guard

    subjects = query_all(
        "SELECT * FROM subjects WHERE status='Active' "
        "ORDER BY subject_name"
    )

    if request.method == "POST":

        values = (
            request.form.get("question_id", "").strip(),
            safe_int(request.form.get("subject_id")),
            request.form.get("question_text", "").strip(),
            request.form.get("option_a", "").strip(),
            request.form.get("option_b", "").strip(),
            request.form.get("option_c", "").strip(),
            request.form.get("option_d", "").strip(),
            normalize_answer(request.form.get("correct_answer")),
            safe_float(request.form.get("marks"), 1),
            safe_float(request.form.get("negative_mark"), 0),
            request.form.get("difficulty", "Medium"),
            request.form.get("topic", "").strip(),
            request.form.get("question_type", "MCQ"),
            request.form.get("status", "Active"),
            request.form.get("explanation", "").strip()
        )

        execute_write("""
            INSERT INTO questions
            (question_id,subject_id,question_text,
             option_a,option_b,option_c,option_d,
             correct_answer,marks,negative_mark,
             difficulty,topic,question_type,status,explanation)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        """, values)

        flash("Question added.")
        return redirect(url_for("admin_questions"))

    body = """
    <div class="card">
        <h1>Add Question</h1>

        <form method="post">

            <div class="row">

                <div>
                    <label>Question ID</label>
                    <input name="question_id"
                           placeholder="Q001">
                </div>

                <div>
                    <label>Subject</label>
                    <select name="subject_id" required>
                        <option value="">Select Subject</option>
                        {% for s in subjects %}
                        <option value="{{ s.id }}">
                            {{ s.subject_code }} -
                            {{ s.subject_name }}
                        </option>
                        {% endfor %}
                    </select>
                </div>

            </div>

            <label>Question</label>
            <textarea name="question_text" required></textarea>

            <div class="row">
                <div>
                    <label>Option A</label>
                    <input name="option_a" required>
                </div>
                <div>
                    <label>Option B</label>
                    <input name="option_b" required>
                </div>
            </div>

            <div class="row">
                <div>
                    <label>Option C</label>
                    <input name="option_c" required>
                </div>
                <div>
                    <label>Option D</label>
                    <input name="option_d" required>
                </div>
            </div>

            <div class="row3">
                <div>
                    <label>Correct Answer</label>
                    <select name="correct_answer">
                        <option>A</option>
                        <option>B</option>
                        <option>C</option>
                        <option>D</option>
                    </select>
                </div>

                <div>
                    <label>Marks</label>
                    <input type="number"
                           step="0.01"
                           min="0"
                           name="marks"
                           value="1">
                </div>

                <div>
                    <label>Negative Mark</label>
                    <input type="number"
                           step="0.01"
                           min="0"
                           name="negative_mark"
                           value="0">
                </div>
            </div>

            <div class="row">
                <div>
                    <label>Difficulty</label>
                    <select name="difficulty">
                        <option>Easy</option>
                        <option selected>Medium</option>
                        <option>Hard</option>
                    </select>
                </div>

                <div>
                    <label>Question Type</label>
                    <select name="question_type">
                        <option>MCQ</option>
                        <option>Multiple Correct</option>
                    </select>
                </div>
            </div>

            <label>Topic</label>
            <input name="topic">

            <label>Status</label>
            <select name="status">
                <option>Active</option>
                <option>Inactive</option>
            </select>

            <label>Explanation</label>
            <textarea name="explanation"></textarea>

            <button class="green">Save Question</button>
            <a class="btn gray"
               href="{{ url_for('admin_questions') }}">
               Cancel
            </a>

        </form>
    </div>
    """

    return page(
        "Add Question",
        body,
        admin=True,
        subjects=subjects
    )


@app.route("/admin/questions/<int:qid>/edit", methods=["GET", "POST"])
def question_edit(qid):

    guard = admin_required()
    if guard:
        return guard

    question = query_one(
        "SELECT * FROM questions WHERE id=?",
        (qid,)
    )

    if not question:
        flash("Question not found.")
        return redirect(url_for("admin_questions"))

    subjects = query_all(
        "SELECT * FROM subjects ORDER BY subject_name"
    )

    if request.method == "POST":

        execute_write("""
            UPDATE questions
            SET question_id=?,
                subject_id=?,
                question_text=?,
                option_a=?,
                option_b=?,
                option_c=?,
                option_d=?,
                correct_answer=?,
                marks=?,
                negative_mark=?,
                difficulty=?,
                topic=?,
                question_type=?,
                status=?,
                explanation=?
            WHERE id=?
        """, (
            request.form.get("question_id", "").strip(),
            safe_int(request.form.get("subject_id")),
            request.form.get("question_text", "").strip(),
            request.form.get("option_a", "").strip(),
            request.form.get("option_b", "").strip(),
            request.form.get("option_c", "").strip(),
            request.form.get("option_d", "").strip(),
            normalize_answer(
                request.form.get("correct_answer")
            ),
            safe_float(request.form.get("marks"), 1),
            safe_float(
                request.form.get("negative_mark"), 0
            ),
            request.form.get("difficulty", "Medium"),
            request.form.get("topic", "").strip(),
            request.form.get("question_type", "MCQ"),
            request.form.get("status", "Active"),
            request.form.get("explanation", "").strip(),
            qid
        ))

        flash("Question updated.")
        return redirect(url_for("admin_questions"))

    body = """
    <div class="card">
        <h1>Edit Question</h1>

        <form method="post">

            <div class="row">
                <div>
                    <label>Question ID</label>
                    <input name="question_id"
                           value="{{ question.question_id }}">
                </div>

                <div>
                    <label>Subject</label>
                    <select name="subject_id" required>
                        {% for s in subjects %}
                        <option value="{{ s.id }}"
                        {% if s.id == question.subject_id %}
                        selected
                        {% endif %}>
                            {{ s.subject_code }} -
                            {{ s.subject_name }}
                        </option>
                        {% endfor %}
                    </select>
                </div>
            </div>

            <label>Question</label>
            <textarea name="question_text"
                      required>{{ question.question_text }}</textarea>

            <div class="row">
                <div>
                    <label>Option A</label>
                    <input name="option_a"
                           value="{{ question.option_a }}"
                           required>
                </div>
                <div>
                    <label>Option B</label>
                    <input name="option_b"
                           value="{{ question.option_b }}"
                           required>
                </div>
            </div>

            <div class="row">
                <div>
                    <label>Option C</label>
                    <input name="option_c"
                           value="{{ question.option_c }}"
                           required>
                </div>
                <div>
                    <label>Option D</label>
                    <input name="option_d"
                           value="{{ question.option_d }}"
                           required>
                </div>
            </div>

            <div class="row3">
                <div>
                    <label>Correct Answer</label>
                    <select name="correct_answer">
                        {% for x in ["A","B","C","D"] %}
                        <option
                        {% if question.correct_answer == x %}
                        selected
                        {% endif %}>
                        {{ x }}
                        </option>
                        {% endfor %}
                    </select>
                </div>

                <div>
                    <label>Marks</label>
                    <input type="number"
                           step="0.01"
                           min="0"
                           name="marks"
                           value="{{ question.marks }}">
                </div>

                <div>
                    <label>Negative Mark</label>
                    <input type="number"
                           step="0.01"
                           min="0"
                           name="negative_mark"
                           value="{{ question.negative_mark }}">
                </div>
            </div>

            <div class="row">
                <div>
                    <label>Difficulty</label>
                    <select name="difficulty">
                        {% for x in ["Easy","Medium","Hard"] %}
                        <option
                        {% if question.difficulty == x %}
                        selected
                        {% endif %}>
                        {{ x }}
                        </option>
                        {% endfor %}
                    </select>
                </div>

                <div>
                    <label>Question Type</label>
                    <select name="question_type">
                        {% for x in ["MCQ","Multiple Correct"] %}
                        <option
                        {% if question.question_type == x %}
                        selected
                        {% endif %}>
                        {{ x }}
                        </option>
                        {% endfor %}
                    </select>
                </div>
            </div>

            <label>Topic</label>
            <input name="topic"
                   value="{{ question.topic }}">

            <label>Status</label>
            <select name="status">
                <option
                {% if question.status == "Active" %}
                selected
                {% endif %}>Active</option>

                <option
                {% if question.status == "Inactive" %}
                selected
                {% endif %}>Inactive</option>
            </select>

            <label>Explanation</label>
            <textarea name="explanation">{{ question.explanation }}</textarea>

            <button class="green">Save Changes</button>

            <a class="btn gray"
               href="{{ url_for('admin_questions') }}">
               Cancel
            </a>
        </form>
    </div>
    """

    return page(
        "Edit Question",
        body,
        admin=True,
        question=question,
        subjects=subjects
    )


@app.route("/admin/questions/<int:qid>/toggle", methods=["POST"])
def question_toggle(qid):

    guard = admin_required()
    if guard:
        return guard

    row = query_one(
        "SELECT status FROM questions WHERE id=?",
        (qid,)
    )

    if row:
        status = (
            "Inactive"
            if row["status"] == "Active"
            else "Active"
        )

        execute_write(
            "UPDATE questions SET status=? WHERE id=?",
            (status, qid)
        )

    return redirect(url_for("admin_questions"))


@app.route("/admin/questions/<int:qid>/delete", methods=["POST"])
def question_delete(qid):

    guard = admin_required()
    if guard:
        return guard

    # Remove dependent records first. Older WTW databases may contain
    # exam_questions even though newer local schemas do not.
    with db_conn() as conn:
        if table_exists(conn, "exam_questions"):
            conn.execute(
                "DELETE FROM exam_questions WHERE question_id=?",
                (qid,)
            )
        if table_exists(conn, "exam_answers"):
            conn.execute(
                "DELETE FROM exam_answers WHERE question_id=?",
                (qid,)
            )
        conn.execute(
            "DELETE FROM questions WHERE id=?",
            (qid,)
        )

    flash("Question deleted successfully.")
    return redirect(url_for("admin_questions"))


@app.route("/admin/questions/delete-selected", methods=["POST"])
def question_delete_selected():

    guard = admin_required()
    if guard:
        return guard

    ids = request.form.getlist("question_ids")

    if not ids:
        flash("No questions selected.")
        return redirect(url_for("admin_questions"))

    for value in ids:
        qid = safe_int(value)
        if qid:
            with db_conn() as conn:
                if table_exists(conn, "exam_questions"):
                    conn.execute(
                        "DELETE FROM exam_questions WHERE question_id=?",
                        (qid,)
                    )
                if table_exists(conn, "exam_answers"):
                    conn.execute(
                        "DELETE FROM exam_answers WHERE question_id=?",
                        (qid,)
                    )
                conn.execute(
                    "DELETE FROM questions WHERE id=?",
                    (qid,)
                )

    flash("Selected questions deleted.")
    return redirect(url_for("admin_questions"))


@app.route("/admin/questions/delete-all", methods=["POST"])
def question_delete_all():

    guard = admin_required()
    if guard:
        return guard

    with db_conn() as conn:
        if table_exists(conn, "exam_questions"):
            conn.execute("DELETE FROM exam_questions")
        if table_exists(conn, "exam_answers"):
            conn.execute("DELETE FROM exam_answers")
        conn.execute("DELETE FROM questions")

    flash("All questions deleted.")
    return redirect(url_for("admin_questions"))


@app.route("/admin/questions/template")
def question_template():

    guard = admin_required()
    if guard:
        return guard

    if Workbook is None:
        flash("openpyxl is not installed.")
        return redirect(url_for("admin_questions"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"

    headers = [
        "Question ID",
        "Subject Code",
        "Question",
        "Option A",
        "Option B",
        "Option C",
        "Option D",
        "Correct Answer"
    ]

    ws.append(headers)

    path = os.path.join(
        UPLOAD_DIR,
        "WTW_Academy_Question_Template.xlsx"
    )

    wb.save(path)

    return send_file(
        path,
        as_attachment=True,
        download_name="WTW_Academy_Question_Template.xlsx"
    )


@app.route("/admin/questions/import", methods=["GET", "POST"])
def question_import():

    guard = admin_required()
    if guard:
        return guard

    if request.method == "POST":

        file = request.files.get("file")

        if not file or not file.filename:
            flash("Please select an Excel file.")
            return redirect(url_for("question_import"))

        if load_workbook is None:
            flash("openpyxl is not installed.")
            return redirect(url_for("question_import"))

        path = os.path.join(
            UPLOAD_DIR,
            "question_import_" +
            str(int(time.time())) +
            ".xlsx"
        )

        file.save(path)

        try:
            wb = load_workbook(
                path,
                read_only=True,
                data_only=True
            )

            ws = wb.active

            rows = list(ws.iter_rows(values_only=True))

            if not rows:
                flash("Excel file is empty.")
                return redirect(url_for("question_import"))

            header = [
                str(x or "").strip().lower()
                for x in rows[0]
            ]

            expected = [
                "question id",
                "subject code",
                "question",
                "option a",
                "option b",
                "option c",
                "option d",
                "correct answer"
            ]

            if header[:8] != expected:
                flash(
                    "Excel format must be: "
                    "Question ID, Subject Code, Question, "
                    "Option A, Option B, Option C, Option D, "
                    "Correct Answer"
                )
                return redirect(url_for("question_import"))

            count = 0

            for row in rows[1:]:

                if not row:
                    continue

                values = list(row) + [""] * 8

                qid = str(values[0] or "").strip()
                code = str(values[1] or "").strip().upper()
                question_text = str(values[2] or "").strip()

                if not question_text:
                    continue

                subject = query_one(
                    "SELECT id FROM subjects "
                    "WHERE UPPER(subject_code)=?",
                    (code,)
                )

                if not subject:
                    continue

                correct = normalize_answer(values[7])

                execute_write("""
                    INSERT INTO questions
                    (question_id,subject_id,question_text,
                     option_a,option_b,option_c,option_d,
                     correct_answer,marks,negative_mark,
                     difficulty,topic,question_type,
                     status,explanation)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """, (
                    qid,
                    subject["id"],
                    question_text,
                    str(values[3] or ""),
                    str(values[4] or ""),
                    str(values[5] or ""),
                    str(values[6] or ""),
                    correct,
                    1,
                    0,
                    "Medium",
                    "",
                    "MCQ",
                    "Active",
                    ""
                ))

                count += 1

            flash(
                str(count) +
                " questions imported successfully."
            )

        except Exception as e:
            flash("Excel import error: " + str(e))

        return redirect(url_for("admin_questions"))

    body = """
    <div class="card">
        <h1>Import Questions</h1>

        <p>
            Excel columns must be exactly:
        </p>

        <p class="badge">
            Question ID | Subject Code | Question |
            Option A | Option B | Option C | Option D |
            Correct Answer
        </p>

        <form method="post"
              enctype="multipart/form-data">

            <label>Excel File</label>
            <input type="file"
                   name="file"
                   accept=".xlsx"
                   required>

            <button class="green">
                Import Questions
            </button>

            <a class="btn gray"
               href="{{ url_for('admin_questions') }}">
               Cancel
            </a>
        </form>
    </div>
    """

    return page(
        "Import Questions",
        body,
        admin=True
    )


# ============================================================
# BATCHES
# ============================================================

@app.route("/admin/batches")
def admin_batches():

    guard = admin_required()
    if guard:
        return guard

    batches = query_all("""
        SELECT b.*,
        (SELECT COUNT(*)
         FROM batch_students bs
         WHERE bs.batch_id=b.id) AS student_count
        FROM batches b
        ORDER BY b.id DESC
    """)

    body = """
    <div class="card">
        <h1>Batches</h1>

        <a class="btn green"
           href="{{ url_for('batch_add') }}">
           Create Batch
        </a>

        <div class="table-scroll">
        <table>
        <tr>
            <th>Batch</th>
            <th>Description</th>
            <th>Students</th>
            <th>Status</th>
            <th>Actions</th>
        </tr>

        {% for b in batches %}
        <tr>
            <td>{{ b.batch_name }}</td>
            <td>{{ b.description }}</td>
            <td>{{ b.student_count }}</td>
            <td>{{ b.status }}</td>
            <td class="actions">

                <a class="btn small"
                   href="{{ url_for(
                       'batch_students',
                       bid=b.id
                   ) }}">
                   Manage Students
                </a>

                <a class="btn small"
                   href="{{ url_for(
                       'batch_edit',
                       bid=b.id
                   ) }}">
                   Edit
                </a>

                <form method="post"
                      action="{{ url_for(
                          'batch_toggle',
                          bid=b.id
                      ) }}"
                      style="display:inline">
                    <button class="btn small gray">
                        Toggle
                    </button>
                </form>

                <form method="post"
                      action="{{ url_for(
                          'batch_delete',
                          bid=b.id
                      ) }}"
                      style="display:inline"
                      onsubmit="return confirm('Delete batch?')">
                    <button class="btn small red">
                        Delete
                    </button>
                </form>

            </td>
        </tr>
        {% endfor %}
        </table>
        </div>
    </div>
    """

    return page(
        "Batches",
        body,
        admin=True,
        batches=batches
    )


@app.route("/admin/batches/add", methods=["GET", "POST"])
def batch_add():

    guard = admin_required()
    if guard:
        return guard

    if request.method == "POST":

        name = request.form.get("batch_name", "").strip()
        description = request.form.get(
            "description", ""
        ).strip()

        batch_code = request.form.get("batch_code", "").strip()

        if not name:
            flash("Batch name is required.", "error")
        else:
            if not batch_code:
                batch_code = "BATCH-" + secrets.token_hex(3).upper()

            try:
                # Keep batch creation compatible with both the current
                # schema and older local databases that may not yet have
                # batch_code.
                execute_write("""
                    INSERT INTO batches
                    (batch_name, name, description, status)
                    VALUES (?,?,?,?)
                """, (
                    name,
                    name,
                    description,
                    "Active"
                ))

                # Add the generated code when the column exists.
                try:
                    latest = query_one("""
                        SELECT id FROM batches
                        ORDER BY id DESC
                        LIMIT 1
                    """)
                    if latest:
                        execute_write("""
                            UPDATE batches
                            SET batch_code=?
                            WHERE id=?
                        """, (
                            batch_code,
                            latest["id"]
                        ))
                except Exception:
                    # Older databases without batch_code can still create
                    # the batch successfully.
                    pass

                flash("Batch created successfully.")
                return redirect(url_for("admin_batches"))
            except Exception as e:
                flash("Batch could not be created: " + str(e), "error")

    body = """
    <div class="card">
        <h1>Create Batch</h1>

        <form method="post">
            <label>Batch Name</label>
            <input name="batch_name" required>

            <label>Batch Code (optional)</label>
            <input name="batch_code" placeholder="Auto-generated if blank">

            <label>Description</label>
            <textarea name="description"></textarea>

            <button class="green">Create Batch</button>
            <a class="btn gray"
               href="{{ url_for('admin_batches') }}">
               Cancel
            </a>
        </form>
    </div>
    """

    return page(
        "Create Batch",
        body,
        admin=True
    )


@app.route("/admin/batches/<int:bid>/edit", methods=["GET", "POST"])
def batch_edit(bid):

    guard = admin_required()
    if guard:
        return guard

    batch = query_one(
        "SELECT * FROM batches WHERE id=?",
        (bid,)
    )

    if not batch:
        flash("Batch not found.")
        return redirect(url_for("admin_batches"))

    if request.method == "POST":

        name = request.form.get("batch_name", "").strip()
        batch_code = request.form.get("batch_code", "").strip()
        description = request.form.get(
            "description", ""
        ).strip()

        if not name:
            flash("Batch name is required.", "error")
            return redirect(url_for("batch_edit", bid=bid))

        if not batch_code:
            batch_code = "BATCH-" + str(bid)

        execute_write("""
            UPDATE batches
            SET batch_name=?,
                batch_code=?,
                name=?,
                description=?
            WHERE id=?
        """, (
            name,
            batch_code,
            name,
            description,
            bid
        ))

        flash("Batch updated.")
        return redirect(url_for("admin_batches"))

    body = """
    <div class="card">
        <h1>Edit Batch</h1>

        <form method="post">
            <label>Batch Name</label>
            <input name="batch_name"
                   value="{{ batch.batch_name }}"
                   required>

            <label>Batch Code</label>
            <input name="batch_code"
                   value="{{ batch.batch_code or ('BATCH-' ~ batch.id) }}">

            <label>Description</label>
            <textarea name="description">{{ batch.description }}</textarea>

            <button class="green">Save</button>
            <a class="btn gray"
               href="{{ url_for('admin_batches') }}">
               Cancel
            </a>
        </form>
    </div>
    """

    return page(
        "Edit Batch",
        body,
        admin=True,
        batch=batch
    )


@app.route("/admin/batches/<int:bid>/toggle", methods=["POST"])
def batch_toggle(bid):

    guard = admin_required()
    if guard:
        return guard

    row = query_one(
        "SELECT status FROM batches WHERE id=?",
        (bid,)
    )

    if row:
        status = (
            "Inactive"
            if row["status"] == "Active"
            else "Active"
        )

        execute_write(
            "UPDATE batches SET status=? WHERE id=?",
            (status, bid)
        )

    return redirect(url_for("admin_batches"))


@app.route("/admin/batches/<int:bid>/delete", methods=["POST"])
def batch_delete(bid):

    guard = admin_required()
    if guard:
        return guard

    execute_write(
        "DELETE FROM batch_students WHERE batch_id=?",
        (bid,)
    )

    execute_write(
        "DELETE FROM exam_batches WHERE batch_id=?",
        (bid,)
    )

    execute_write(
        "DELETE FROM batches WHERE id=?",
        (bid,)
    )

    flash("Batch deleted.")
    return redirect(url_for("admin_batches"))


@app.route("/admin/batches/<int:bid>/students",
           methods=["GET", "POST"])
def batch_students(bid):

    guard = admin_required()
    if guard:
        return guard

    batch = query_one(
        "SELECT * FROM batches WHERE id=?",
        (bid,)
    )

    if not batch:
        flash("Batch not found.")
        return redirect(url_for("admin_batches"))

    if request.method == "POST":

        selected = request.form.getlist("student_ids")

        execute_write(
            "DELETE FROM batch_students WHERE batch_id=?",
            (bid,)
        )

        for sid in selected:
            sid_int = safe_int(sid)
            if sid_int:
                execute_write("""
                    INSERT OR IGNORE INTO batch_students
                    (batch_id,student_id)
                    VALUES (?,?)
                """, (
                    bid,
                    sid_int
                ))

        flash("Batch students updated.")
        return redirect(
            url_for("batch_students", bid=bid)
        )

    students = query_all("""
        SELECT *
        FROM students
        WHERE status='Active'
        ORDER BY student_name,name
    """)

    selected_rows = query_all(
        "SELECT student_id FROM batch_students "
        "WHERE batch_id=?",
        (bid,)
    )

    selected_ids = {
        row["student_id"]
        for row in selected_rows
    }

    body = """
    <div class="card">
        <h1>Manage Batch Students</h1>
        <p><b>{{ batch.batch_name }}</b></p>

        <form method="post">

        {% for s in students %}
        <label style="font-weight:400">
            <input type="checkbox"
                   name="student_ids"
                   value="{{ s.id }}"
                   style="width:auto"
                   {% if s.id in selected_ids %}
                   checked
                   {% endif %}>
            {{ s.student_id }} -
            {{ s.student_name or s.name }}
        </label>
        {% endfor %}

        <button class="green">
            Save Selected Students
        </button>

        <a class="btn gray"
           href="{{ url_for('admin_batches') }}">
           Back
        </a>

        </form>
    </div>
    """

    return page(
        "Batch Students",
        body,
        admin=True,
        batch=batch,
        students=students,
        selected_ids=selected_ids
    )


# ============================================================
# EXAMS
# ============================================================

@app.route("/admin/exams")
def admin_exams():

    guard = admin_required()
    if guard:
        return guard

    exams = query_all("""
        SELECT e.*, s.subject_name, s.subject_code
        FROM exams e
        LEFT JOIN subjects s
        ON e.subject_id=s.id
        ORDER BY e.id DESC
    """)

    body = """
    <div class="card">
        <h1>Exams</h1>

        <a class="btn green"
           href="{{ url_for('exam_add') }}">
           Create Exam
        </a>

        <form method="post"
              action="{{ url_for('exam_delete_selected') }}"
              style="display:inline"
              onsubmit="return confirm('Delete selected exams?')">
            <button class="btn red">
                Delete Selected
            </button>

            {% for e in exams %}
            <input type="hidden"
                   class="selected-exam-holder"
                   name="exam_ids"
                   value="{{ e.id }}"
                   disabled>
            {% endfor %}
        </form>

        <form method="post"
              action="{{ url_for('exam_delete_all') }}"
              style="display:inline"
              onsubmit="return confirm('DELETE ALL EXAMS?')">
            <button class="btn red">
                Delete All Exams
            </button>
        </form>
    </div>

    <div class="card">
        <div class="table-scroll">
        <table>
        <tr>
            <th>✓</th>
            <th>Exam</th>
            <th>Subject</th>
            <th>Questions</th>
            <th>Marks</th>
            <th>Time</th>
            <th>Attempts</th>
            <th>Fee</th>
            <th>Status</th>
            <th>Action</th>
        </tr>

        {% for e in exams %}
        <tr>
            <td>
                <input type="checkbox"
                       class="exambox"
                       data-id="{{ e.id }}">
            </td>

            <td>
                <b>{{ e.exam_name or e.name }}</b>
            </td>

            <td>
                {{ e.subject_code }}
                {{ e.subject_name }}
            </td>

            <td>{{ e.number_of_questions }}</td>
            <td>{{ e.marks_per_question }}</td>
            <td>{{ e.duration_minutes }} min</td>

            <td>
                {% if e.attempts_allowed == 0 %}
                    Unlimited
                {% else %}
                    {{ e.attempts_allowed }}
                {% endif %}
            </td>

            <td>{{ e.fee }}</td>
            <td>{{ e.status }}</td>

            <td class="actions">

                <a class="btn small"
                   href="{{ url_for(
                       'exam_edit',
                       eid=e.id
                   ) }}">
                   Edit
                </a>

                <a class="btn small orange"
                   href="{{ url_for(
                       'exam_assign',
                       eid=e.id
                   ) }}">
                   Assign
                </a>

                <form method="post"
                      action="{{ url_for(
                          'exam_toggle',
                          eid=e.id
                      ) }}"
                      style="display:inline">
                    <button class="btn small gray">
                        Toggle
                    </button>
                </form>

                <form method="post"
                      action="{{ url_for(
                          'exam_delete',
                          eid=e.id
                      ) }}"
                      style="display:inline"
                      onsubmit="return confirm('Delete exam?')">
                    <button class="btn small red">
                        Delete
                    </button>
                </form>

            </td>
        </tr>
        {% endfor %}
        </table>
        </div>
    </div>

    <script>
    const boxes=document.querySelectorAll('.exambox');

    boxes.forEach(function(box){
        box.addEventListener('change',function(){
            const id=this.dataset.id;

            const holders=document.querySelectorAll(
                '.selected-exam-holder'
            );

            holders.forEach(function(h){
                if(h.value===id){
                    h.disabled=!box.checked;
                }
            });
        });
    });
    </script>
    """

    return page(
        "Exams",
        body,
        admin=True,
        exams=exams
    )


@app.route("/admin/exams/add", methods=["GET", "POST"])
def exam_add():

    guard = admin_required()
    if guard:
        return guard

    subjects = query_all(
        "SELECT * FROM subjects WHERE status='Active' "
        "ORDER BY subject_name"
    )

    if request.method == "POST":

        name = request.form.get("exam_name", "").strip()
        exam_id = request.form.get(
            "exam_id", ""
        ).strip()

        subject_id = safe_int(
            request.form.get("subject_id")
        )

        nq = safe_int(
            request.form.get("number_of_questions"),
            0
        )

        marks = safe_float(
            request.form.get("marks_per_question"),
            1
        )

        duration = safe_int(
            request.form.get("duration_minutes"),
            30
        )

        fee = safe_float(
            request.form.get("fee"),
            0
        )

        attempts = safe_int(
            request.form.get("attempts_allowed"),
            1
        )

        if attempts < 0:
            attempts = 1

        if not name or not subject_id:
            flash("Exam Name and Subject are required.")
        else:

            execute_write("""
                INSERT INTO exams
                (exam_id,exam_name,name,subject_id,
                 number_of_questions,marks_per_question,
                 duration_minutes,fee,
                 attempts_allowed,max_attempts,status)
                VALUES (?,?,?,?,?,?,?,?,?,?,?)
            """, (
                exam_id,
                name,
                name,
                subject_id,
                nq,
                marks,
                duration,
                fee,
                attempts,
                attempts,
                "Active"
            ))

            flash("Exam created.")
            return redirect(url_for("admin_exams"))

    body = """
    <div class="card">
        <h1>Create Exam</h1>

        <form method="post">

            <label>Exam ID</label>
            <input name="exam_id"
                   placeholder="EXAM001">

            <label>Exam Name</label>
            <input name="exam_name" required>

            <label>Subject</label>
            <select name="subject_id" required>
                <option value="">Select Subject</option>
                {% for s in subjects %}
                <option value="{{ s.id }}">
                    {{ s.subject_code }} -
                    {{ s.subject_name }}
                </option>
                {% endfor %}
            </select>

            <div class="row3">

                <div>
                    <label>Number of Questions</label>
                    <input type="number"
                           min="0"
                           name="number_of_questions"
                           value="0">
                </div>

                <div>
                    <label>Marks per Question</label>
                    <input type="number"
                           step="0.01"
                           min="0"
                           name="marks_per_question"
                           value="1">
                </div>

                <div>
                    <label>Duration - Minutes</label>
                    <input type="number"
                           min="1"
                           name="duration_minutes"
                           value="30">
                </div>

            </div>

            <div class="row">

                <div>
                    <label>Exam Fee</label>
                    <input type="number"
                           step="0.01"
                           min="0"
                           name="fee"
                           value="0">
                </div>

                <div>
                    <label>Attempts Allowed</label>
                    <select name="attempts_allowed">
                        <option value="1">1 Attempt</option>
                        <option value="2">2 Attempts</option>
                        <option value="3">3 Attempts</option>
                        <option value="5">5 Attempts</option>
                        <option value="0">Unlimited</option>
                    </select>
                </div>

            </div>

            <button class="green">Create Exam</button>

            <a class="btn gray"
               href="{{ url_for('admin_exams') }}">
               Cancel
            </a>

        </form>
    </div>
    """

    return page(
        "Create Exam",
        body,
        admin=True,
        subjects=subjects
    )


@app.route("/admin/exams/<int:eid>/edit",
           methods=["GET", "POST"])
def exam_edit(eid):

    guard = admin_required()
    if guard:
        return guard

    exam = query_one(
        "SELECT * FROM exams WHERE id=?",
        (eid,)
    )

    if not exam:
        flash("Exam not found.")
        return redirect(url_for("admin_exams"))

    subjects = query_all(
        "SELECT * FROM subjects ORDER BY subject_name"
    )

    if request.method == "POST":

        attempts = safe_int(
            request.form.get("attempts_allowed"),
            1
        )

        execute_write("""
            UPDATE exams
            SET exam_id=?,
                exam_name=?,
                name=?,
                subject_id=?,
                number_of_questions=?,
                marks_per_question=?,
                duration_minutes=?,
                fee=?,
                attempts_allowed=?,
                max_attempts=?
            WHERE id=?
        """, (
            request.form.get("exam_id", "").strip(),
            request.form.get("exam_name", "").strip(),
            request.form.get("exam_name", "").strip(),
            safe_int(request.form.get("subject_id")),
            safe_int(
                request.form.get("number_of_questions")
            ),
            safe_float(
                request.form.get("marks_per_question"),
                1
            ),
            safe_int(
                request.form.get("duration_minutes"),
                30
            ),
            safe_float(
                request.form.get("fee"),
                0
            ),
            attempts,
            attempts,
            eid
        ))

        flash("Exam updated.")
        return redirect(url_for("admin_exams"))

    body = """
    <div class="card">
        <h1>Edit Exam</h1>

        <form method="post">

            <label>Exam ID</label>
            <input name="exam_id"
                   value="{{ exam.exam_id }}">

            <label>Exam Name</label>
            <input name="exam_name"
                   value="{{ exam.exam_name or exam.name }}"
                   required>

            <label>Subject</label>
            <select name="subject_id" required>
                {% for s in subjects %}
                <option value="{{ s.id }}"
                {% if s.id == exam.subject_id %}
                selected
                {% endif %}>
                    {{ s.subject_code }} -
                    {{ s.subject_name }}
                </option>
                {% endfor %}
            </select>

            <div class="row3">
                <div>
                    <label>Number of Questions</label>
                    <input type="number"
                           min="0"
                           name="number_of_questions"
                           value="{{ exam.number_of_questions }}">
                </div>

                <div>
                    <label>Marks per Question</label>
                    <input type="number"
                           step="0.01"
                           min="0"
                           name="marks_per_question"
                           value="{{ exam.marks_per_question }}">
                </div>

                <div>
                    <label>Duration Minutes</label>
                    <input type="number"
                           min="1"
                           name="duration_minutes"
                           value="{{ exam.duration_minutes }}">
                </div>
            </div>

            <div class="row">
                <div>
                    <label>Fee</label>
                    <input type="number"
                           step="0.01"
                           min="0"
                           name="fee"
                           value="{{ exam.fee }}">
                </div>

                <div>
                    <label>Attempts Allowed</label>
                    <select name="attempts_allowed">
                        {% for x in [1,2,3,5,0] %}
                        <option value="{{ x }}"
                        {% if exam.attempts_allowed == x %}
                        selected
                        {% endif %}>
                            {% if x == 0 %}
                            Unlimited
                            {% else %}
                            {{ x }}
                            {% endif %}
                        </option>
                        {% endfor %}
                    </select>
                </div>
            </div>

            <button class="green">Save</button>

            <a class="btn gray"
               href="{{ url_for('admin_exams') }}">
               Cancel
            </a>
        </form>
    </div>
    """

    return page(
        "Edit Exam",
        body,
        admin=True,
        exam=exam,
        subjects=subjects
    )


@app.route("/admin/exams/<int:eid>/toggle", methods=["POST"])
def exam_toggle(eid):

    guard = admin_required()
    if guard:
        return guard

    row = query_one(
        "SELECT status FROM exams WHERE id=?",
        (eid,)
    )

    if row:
        status = (
            "Inactive"
            if row["status"] == "Active"
            else "Active"
        )

        execute_write(
            "UPDATE exams SET status=? WHERE id=?",
            (status, eid)
        )

    return redirect(url_for("admin_exams"))


def delete_exam_data(eid):

    attempts = query_all(
        "SELECT id FROM exam_attempts WHERE exam_id=?",
        (eid,)
    )

    for a in attempts:
        execute_write(
            "DELETE FROM exam_answers WHERE attempt_id=?",
            (a["id"],)
        )

    execute_write(
        "DELETE FROM exam_attempts WHERE exam_id=?",
        (eid,)
    )

    execute_write(
        "DELETE FROM exam_students WHERE exam_id=?",
        (eid,)
    )

    execute_write(
        "DELETE FROM exam_batches WHERE exam_id=?",
        (eid,)
    )

    execute_write(
        "DELETE FROM payments WHERE exam_id=?",
        (eid,)
    )

    execute_write(
        "DELETE FROM exams WHERE id=?",
        (eid,)
    )


@app.route("/admin/exams/<int:eid>/delete",
           methods=["POST"])
def exam_delete(eid):

    guard = admin_required()
    if guard:
        return guard

    delete_exam_data(eid)

    flash("Exam deleted.")
    return redirect(url_for("admin_exams"))


@app.route("/admin/exams/delete-selected",
           methods=["POST"])
def exam_delete_selected():

    guard = admin_required()
    if guard:
        return guard

    ids = request.form.getlist("exam_ids")

    for value in ids:
        eid = safe_int(value)
        if eid:
            delete_exam_data(eid)

    flash("Selected exams deleted.")
    return redirect(url_for("admin_exams"))


@app.route("/admin/exams/delete-all",
           methods=["POST"])
def exam_delete_all():

    guard = admin_required()
    if guard:
        return guard

    exams = query_all(
        "SELECT id FROM exams"
    )

    for e in exams:
        delete_exam_data(e["id"])

    flash("All exams deleted.")
    return redirect(url_for("admin_exams"))


# ============================================================
# EXAM ASSIGNMENT
# ============================================================

@app.route("/admin/exams/<int:eid>/assign",
           methods=["GET", "POST"])
def exam_assign(eid):

    guard = admin_required()
    if guard:
        return guard

    exam = query_one(
        "SELECT * FROM exams WHERE id=?",
        (eid,)
    )

    if not exam:
        flash("Exam not found.")
        return redirect(url_for("admin_exams"))

    batches = query_all("""
        SELECT *
        FROM batches
        WHERE status='Active'
        ORDER BY batch_name
    """)

    students = query_all("""
        SELECT *
        FROM students
        WHERE status='Active'
        ORDER BY student_name,name
    """)

    if request.method == "POST":

        assign_type = request.form.get(
            "assign_type", "student"
        )

        if assign_type == "batch":

            batch_ids = request.form.getlist(
                "batch_ids"
            )

            for value in batch_ids:

                bid = safe_int(value)

                if not bid:
                    continue

                execute_write("""
                    INSERT OR IGNORE INTO exam_batches
                    (exam_id,batch_id,payment_required,assigned_by)
                    VALUES (?,?,?,?)
                """, (
                    eid,
                    bid,
                    1 if safe_float(exam["fee"]) > 0 else 0,
                    session.get("admin_id", 0)
                ))

            flash("Exam assigned to selected batches.")

        else:

            student_ids = request.form.getlist(
                "student_ids"
            )

            for value in student_ids:

                sid = safe_int(value)

                if not sid:
                    continue

                execute_write("""
                    INSERT OR IGNORE INTO exam_students
                    (exam_id,student_id,payment_required,
                     payment_status,assigned_by)
                    VALUES (?,?,?,?,?)
                """, (
                    eid,
                    sid,
                    1 if safe_float(exam["fee"]) > 0 else 0,
                    "Pending"
                    if safe_float(exam["fee"]) > 0
                    else "Not Required",
                    session.get("admin_id", 0)
                ))

            flash("Exam assigned to selected students.")

        return redirect(
            url_for("exam_assign", eid=eid)
        )

    assigned_students = query_all("""
        SELECT s.student_id,
               s.student_name,
               s.name
        FROM exam_students es
        JOIN students s
        ON s.id=es.student_id
        WHERE es.exam_id=?
        ORDER BY s.student_name,s.name
    """, (eid,))

    assigned_batches = query_all("""
        SELECT b.batch_name
        FROM exam_batches eb
        JOIN batches b
        ON b.id=eb.batch_id
        WHERE eb.exam_id=?
        ORDER BY b.batch_name
    """, (eid,))

    body = """
    <div class="card">

        <h1>Assign Exam</h1>

        <p>
            <b>{{ exam.exam_name or exam.name }}</b>
        </p>

        <form method="post">

            <label>Assignment Type</label>

            <select name="assign_type"
                    id="assign_type"
                    onchange="switchAssign()">
                <option value="student">
                    Individual Students
                </option>
                <option value="batch">
                    Batches
                </option>
            </select>

            <div id="studentBox">

                <label>Select Students</label>

                {% for s in students %}
                <label style="font-weight:400">
                    <input type="checkbox"
                           name="student_ids"
                           value="{{ s.id }}"
                           style="width:auto">
                    {{ s.student_id }} -
                    {{ s.student_name or s.name }}
                </label>
                {% endfor %}

            </div>

            <div id="batchBox" style="display:none">

                <label>Select Batches</label>

                {% for b in batches %}
                <label style="font-weight:400">
                    <input type="checkbox"
                           name="batch_ids"
                           value="{{ b.id }}"
                           style="width:auto">
                    {{ b.batch_name }}
                </label>
                {% endfor %}

            </div>

            <button class="green">
                Assign Exam
            </button>

        </form>
    </div>

    <div class="card">
        <h2>Direct Student Assignments</h2>

        {% for s in assigned_students %}
        <div class="badge green">
            {{ s.student_id }} -
            {{ s.student_name or s.name }}
        </div>
        {% else %}
        <p class="muted">No direct students assigned.</p>
        {% endfor %}
    </div>

    <div class="card">
        <h2>Batch Assignments</h2>

        {% for b in assigned_batches %}
        <div class="badge orange">
            {{ b.batch_name }}
        </div>
        {% else %}
        <p class="muted">No batches assigned.</p>
        {% endfor %}
    </div>

    <script>
    function switchAssign(){
        const type=document.getElementById('assign_type').value;

        document.getElementById('studentBox').style.display =
            type==='student' ? 'block' : 'none';

        document.getElementById('batchBox').style.display =
            type==='batch' ? 'block' : 'none';
    }
    </script>
    """

    return page(
        "Assign Exam",
        body,
        admin=True,
        exam=exam,
        batches=batches,
        students=students,
        assigned_students=assigned_students,
        assigned_batches=assigned_batches
    )


# ============================================================
# PAYMENTS
# ============================================================

@app.route("/admin/payments")
def admin_payments():

    guard = admin_required()
    if guard:
        return guard

    payments = query_all("""
        SELECT p.*,
               e.exam_name,
               s.student_id,
               s.student_name,
               s.name
        FROM payments p
        LEFT JOIN exams e
        ON e.id=p.exam_id
        LEFT JOIN students s
        ON s.id=p.student_id
        ORDER BY p.id DESC
    """)

    body = """
    <div class="card">
        <h1>Payments</h1>

        <div class="table-scroll">
        <table>
        <tr>
            <th>Date</th>
            <th>Student</th>
            <th>Exam</th>
            <th>Amount</th>
            <th>Reference</th>
            <th>Status</th>
            <th>Action</th>
        </tr>

        {% for p in payments %}
        <tr>
            <td>{{ p.created_at }}</td>
            <td>
                {{ p.student_id }} -
                {{ p.student_name or p.name }}
            </td>
            <td>{{ p.exam_name }}</td>
            <td>{{ p.amount }}</td>
            <td>{{ p.upi_reference or p.transaction_id }}</td>
            <td>{{ p.status }}</td>
            <td>
                {% if p.status == "Pending" %}
                <form method="post"
                      action="{{ url_for(
                          'payment_verify',
                          pid=p.id
                      ) }}">
                    <button class="btn small green">
                        Verify
                    </button>
                </form>
                {% endif %}
            </td>
        </tr>
        {% endfor %}
        </table>
        </div>
    </div>
    """

    return page(
        "Payments",
        body,
        admin=True,
        payments=payments
    )


@app.route("/admin/payments/<int:pid>/verify",
           methods=["POST"])
def payment_verify(pid):

    guard = admin_required()
    if guard:
        return guard

    payment = query_one(
        "SELECT * FROM payments WHERE id=?",
        (pid,)
    )

    if payment:

        execute_write("""
            UPDATE payments
            SET status='Verified',
                verified_at=?
            WHERE id=?
        """, (
            now_text(),
            pid
        ))

        execute_write("""
            UPDATE exam_students
            SET payment_status='Verified',
                payment_reference=(
                    SELECT COALESCE(
                        upi_reference,
                        transaction_id
                    )
                    FROM payments
                    WHERE id=?
                ),
                paid_at=?
            WHERE exam_id=?
            AND student_id=?
        """, (
            pid,
            now_text(),
            payment["exam_id"],
            payment["student_id"]
        ))

        flash("Payment verified.")

    return redirect(url_for("admin_payments"))


# ============================================================
# RESULTS - ADMIN
# ============================================================

@app.route("/admin/results")
def admin_results():

    guard = admin_required()
    if guard:
        return guard

    results = query_all("""
        SELECT a.*,
               e.exam_name,
               s.student_id,
               s.student_name,
               s.name
        FROM exam_attempts a
        JOIN exams e
        ON e.id=a.exam_id
        JOIN students s
        ON s.id=a.student_id
        WHERE a.status='Submitted'
        ORDER BY a.id DESC
    """)

    body = """
    <div class="card">
        <h1>Results</h1>

        <div class="table-scroll">
        <table>
        <tr>
            <th>Date</th>
            <th>Student</th>
            <th>Exam</th>
            <th>Attempt</th>
            <th>Score</th>
            <th>Total</th>
        </tr>

        {% for r in results %}
        <tr>
            <td>{{ r.submitted_at }}</td>
            <td>
                {{ r.student_id }} -
                {{ r.student_name or r.name }}
            </td>
            <td>{{ r.exam_name }}</td>
            <td>{{ r.attempt_no }}</td>
            <td>{{ r.score }}</td>
            <td>{{ r.total_marks }}</td>
        </tr>
        {% endfor %}
        </table>
        </div>
    </div>
    """

    return page(
        "Results",
        body,
        admin=True,
        results=results
    )

# ============================================================
# STUDENT LOGIN
# ============================================================

@app.route("/student-login", methods=["GET", "POST"])
def student_login():

    if request.method == "POST":

        student_id = request.form.get(
            "student_id", ""
        ).strip()

        password = request.form.get(
            "password", ""
        )

        student = query_one("""
            SELECT *
            FROM students
            WHERE student_id=?
        """, (student_id,))

        if student and student["status"] == "Active":

            stored = (
                student["password"]
                or student["password_hash"]
            )

            if verify_password(stored, password):

                session.permanent = True
                session["student_db_id"] = student["id"]
                session["student_id"] = student["student_id"]
                session["student_name"] = (
                    student["student_name"]
                    or student["name"]
                )

                return redirect(
                    url_for("student_dashboard")
                )

        flash("Invalid student login.")

    body = """
    <div class="login-wrap">
        <div class="login-card">
            <h1>WTW Academy</h1>
            <h2>Student Login</h2>

            <form method="post">

                <label>Student ID</label>
                <input name="student_id"
                       required
                       autofocus>

                <label>Password</label>

                <div style="position:relative">
                    <input type="password"
                           name="password"
                           id="studentPassword"
                           required
                           style="padding-right:45px">

                    <button type="button"
                            onclick="togglePassword('studentPassword', this)"
                            style="position:absolute;
                                   right:8px;
                                   top:50%;
                                   transform:translateY(-50%);
                                   border:0;
                                   background:none;
                                   cursor:pointer;color:#334155;font-size:20px">
                        👁
                    </button>
                </div>

                <button style="width:100%;margin-top:15px">
                    Login
                </button>

            </form>
        </div>
    </div>
    """

    return page(
        "Student Login",
        body,
        nav=False
    )


@app.route("/student/logout")
def student_logout():

    session.pop("student_db_id", None)
    session.pop("student_id", None)
    session.pop("student_name", None)

    return redirect(url_for("student_login"))

# ============================================================
# STUDENT EXAM ACCESS
# ============================================================

def student_has_exam(student_id, exam_id):

    direct = query_one("""
        SELECT es.*
        FROM exam_students es
        WHERE es.exam_id=?
        AND es.student_id=?
    """, (
        exam_id,
        student_id
    ))

    if direct:
        return direct

    batch = query_one("""
        SELECT eb.*
        FROM exam_batches eb
        JOIN batch_students bs
        ON bs.batch_id=eb.batch_id
        WHERE eb.exam_id=?
        AND bs.student_id=?
        LIMIT 1
    """, (
        exam_id,
        student_id
    ))

    if batch:
        return {
            "payment_required":
                batch["payment_required"],
            "payment_status":
                "Pending"
                if batch["payment_required"]
                else "Not Required"
        }

    return None


def create_direct_assignment_from_batch(
    student_id,
    exam_id,
    assignment
):

    existing = query_one("""
        SELECT *
        FROM exam_students
        WHERE exam_id=?
        AND student_id=?
    """, (
        exam_id,
        student_id
    ))

    if existing:
        return existing

    exam = query_one(
        "SELECT fee FROM exams WHERE id=?",
        (exam_id,)
    )

    fee = safe_float(exam["fee"]) if exam else 0

    execute_write("""
        INSERT OR IGNORE INTO exam_students
        (exam_id,student_id,payment_required,
         payment_status,assigned_by)
        VALUES (?,?,?,?,?)
    """, (
        exam_id,
        student_id,
        1 if fee > 0 else 0,
        "Pending" if fee > 0 else "Not Required",
        0
    ))

    return query_one("""
        SELECT *
        FROM exam_students
        WHERE exam_id=?
        AND student_id=?
    """, (
        exam_id,
        student_id
    ))


def get_exam_questions(exam):

    sql = """
        SELECT *
        FROM questions
        WHERE subject_id=?
        AND status='Active'
        ORDER BY id
    """

    questions = query_all(
        sql,
        (exam["subject_id"],)
    )

    questions = list(questions)

    number = safe_int(
        exam["number_of_questions"],
        0
    )

    if number > 0 and len(questions) > number:
        questions = random.sample(
            questions,
            number
        )

    random.shuffle(questions)

    return questions


# ============================================================
# STUDENT DASHBOARD
# ============================================================

@app.route("/student/dashboard")
def student_dashboard():

    guard = student_required()
    if guard:
        return guard

    sid = session["student_db_id"]

    direct = query_all("""
        SELECT e.*,
               es.payment_required,
               es.payment_status
        FROM exam_students es
        JOIN exams e
        ON e.id=es.exam_id
        WHERE es.student_id=?
        AND e.status='Active'
    """, (sid,))

    batch = query_all("""
        SELECT DISTINCT e.*,
               eb.payment_required,
               CASE
                   WHEN eb.payment_required=1
                   THEN 'Pending'
                   ELSE 'Not Required'
               END AS payment_status
        FROM exam_batches eb
        JOIN batch_students bs
        ON bs.batch_id=eb.batch_id
        JOIN exams e
        ON e.id=eb.exam_id
        WHERE bs.student_id=?
        AND e.status='Active'
    """, (sid,))

    exams = {}
    for e in direct:
        exams[e["id"]] = e

    for e in batch:
        if e["id"] not in exams:
            exams[e["id"]] = e

    exams = list(exams.values())

    exam_data = []

    for e in exams:

        attempts = query_all("""
            SELECT *
            FROM exam_attempts
            WHERE exam_id=?
            AND student_id=?
            ORDER BY attempt_no DESC
        """, (
            e["id"],
            sid
        ))

        in_progress = next(
            (
                a for a in attempts
                if a["status"] == "In Progress"
            ),
            None
        )

        submitted = len([
            a for a in attempts
            if a["status"] == "Submitted"
        ])

        limit = safe_int(
            e["attempts_allowed"],
            1
        )

        remaining = (
            "Unlimited"
            if limit == 0
            else max(0, limit - submitted)
        )

        exam_data.append({
            "exam": e,
            "in_progress": in_progress,
            "submitted": submitted,
            "remaining": remaining
        })

    body = """
    <div class="card">
        <h1>Student Dashboard</h1>
        <p>
            Welcome,
            <b>{{ student_name }}</b>
        </p>
    </div>

    {% for item in exam_data %}

    <div class="card">

        <h2>{{ item.exam.exam_name or item.exam.name }}</h2>

        <p>
            Duration:
            {{ item.exam.duration_minutes }} minutes
            |
            Attempts:
            {% if item.exam.attempts_allowed == 0 %}
            Unlimited
            {% else %}
            {{ item.exam.attempts_allowed }}
            {% endif %}
        </p>

        <p>
            Remaining Attempts:
            <b>{{ item.remaining }}</b>
        </p>

        {% if item.exam.fee|float > 0 %}
        <p>
            Fee:
            <b>{{ item.exam.fee }}</b>
        </p>
        {% endif %}

        {% if item.in_progress %}

        <a class="btn green"
           href="{{ url_for(
               'student_take',
               attempt_id=item.in_progress.id
           ) }}">
           Resume Exam
        </a>

        {% elif item.remaining == 0 %}

        <span class="badge red">
            Attempt Limit Reached
        </span>

        {% else %}

                <a class="btn {% if item.exam.fee|float > 0 and item.exam.payment_status != 'Verified' %}red{% else %}green{% endif %}"
           href="{{ url_for(
               'student_start_exam',
               exam_id=item.exam.id
           ) }}">
           Start Exam
        </a>

        {% endif %}

    </div>

    {% else %}

    <div class="card center">
        <h2>No exams assigned</h2>
        <p class="muted">
            Please contact WTW Academy.
        </p>
    </div>

    {% endfor %}
    """

    return page(
        "Student Dashboard",
        body,
        student=True,
        exam_data=exam_data,
        student_name=session.get("student_name", "")
    )


# ============================================================
# START EXAM
# ============================================================

@app.route("/student/exam/<int:exam_id>/start")
@app.route("/student/exam/<int:exam_id>")
def student_start_exam(exam_id):

    guard = student_required()
    if guard:
        return guard

    sid = session["student_db_id"]

    exam = query_one("""
        SELECT *
        FROM exams
        WHERE id=?
        AND status='Active'
    """, (exam_id,))

    if not exam:
        flash("Exam not available.")
        return redirect(url_for("student_dashboard"))

    assignment = student_has_exam(
        sid,
        exam_id
    )

    if not assignment:
        flash("This exam is not assigned to you.")
        return redirect(url_for("student_dashboard"))

    # Batch assignment becomes direct assignment
    if not query_one("""
        SELECT id
        FROM exam_students
        WHERE exam_id=?
        AND student_id=?
    """, (exam_id, sid)):
        assignment = create_direct_assignment_from_batch(
            sid,
            exam_id,
            assignment
        )

    # Resume existing attempt
    current = query_one("""
        SELECT *
        FROM exam_attempts
        WHERE exam_id=?
        AND student_id=?
        AND status='In Progress'
        ORDER BY id DESC
        LIMIT 1
    """, (
        exam_id,
        sid
    ))

    if current:
        return redirect(
            url_for(
                "student_take",
                attempt_id=current["id"]
            )
        )

    attempts = query_all("""
        SELECT *
        FROM exam_attempts
        WHERE exam_id=?
        AND student_id=?
        AND status='Submitted'
    """, (
        exam_id,
        sid
    ))

    limit = safe_int(
        exam["attempts_allowed"],
        1
    )

    if limit > 0 and len(attempts) >= limit:
        flash("Attempt limit reached.")
        return redirect(url_for("student_dashboard"))

    payment_required = (
        safe_int(assignment["payment_required"], 0)
        == 1
    )

    payment_status = assignment["payment_status"]

    if payment_required and payment_status != "Verified":
        return redirect(
            url_for(
                "student_payment",
                exam_id=exam_id
            )
        )

    questions = get_exam_questions(exam)

    if not questions:
        flash("No active questions available for this exam.")
        return redirect(url_for("student_dashboard"))

    question_order = ",".join(
        str(q["id"])
        for q in questions
    )

    attempt_no = len(attempts) + 1

    attempt_id = execute_write("""
    INSERT INTO exam_attempts
    (exam_id, student_id, attempt_no, status,
     question_order, current_index,
     current_question, last_question,
     started_at, elapsed_seconds,
     score, total_marks)
    VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
""", (
    exam_id,
    sid,
    attempt_no,
    "In Progress",
    question_order,
    0,
    questions[0]["id"],
    questions[-1]["id"],
    now_text(),
    0,
    0,
    sum(
        safe_float(q["marks"], 1)
        for q in questions
    )
))

    return redirect(
        url_for(
            "student_take",
            attempt_id=attempt_id
        )
    )




# ============================================================
# PAYMENT PAGE
# ============================================================

@app.route("/student/payment/<int:exam_id>",
           methods=["GET", "POST"])
def student_payment(exam_id):

    guard = student_required()
    if guard:
        return guard

    sid = session["student_db_id"]

    exam = query_one(
        "SELECT * FROM exams WHERE id=?",
        (exam_id,)
    )

    if not exam:
        flash("Exam not found.")
        return redirect(url_for("student_dashboard"))

    if request.method == "POST":

        reference = request.form.get(
            "upi_reference",
            ""
        ).strip()

        if not reference:
            flash("Enter payment reference.")

        else:

            payment = query_one("""
                SELECT id
                FROM payments
                WHERE exam_id=?
                AND student_id=?
                AND status='Pending'
                ORDER BY id DESC
                LIMIT 1
            """, (
                exam_id,
                sid
            ))

            if payment:

                execute_write("""
                    UPDATE payments
                    SET upi_reference=?
                    WHERE id=?
                """, (
                    reference,
                    payment["id"]
                ))

            else:

                execute_write("""
                    INSERT INTO payments
                    (
                        exam_id,
                        student_id,
                        amount,
                        upi_reference,
                        status
                    )
                    VALUES (?,?,?,?,?)
                """, (
                    exam_id,
                    sid,
                    safe_float(exam["fee"]),
                    reference,
                    "Pending"
                ))

            flash(
                "Payment submitted. "
                "Please wait for admin verification."
            )

            # Payment is Pending.
            # Student returns to dashboard.
            return redirect(
                url_for("student_dashboard")
            )

    # WTW Academy UPI ID
    upi_id = "kavitha94459@oksbi"

    body = """
    <div class="card center">

        <h1>Exam Payment</h1>

        <h2>
            {{ exam.exam_name or exam.name }}
        </h2>

        <p>
            Amount:
            <b>₹{{ exam.fee }}</b>
        </p>

        <p>
            UPI ID:
            <b>{{ upi_id }}</b>
        </p>

        <a class="btn orange"
           href="{{ url_for(
               'student_payment_qr',
               exam_id=exam.id
           ) }}">
            Show QR
        </a>

        <form method="post">

            <label style="text-align:left">
                UPI Transaction Reference
            </label>

            <input
                name="upi_reference"
                placeholder="Enter UPI transaction reference"
                required
            >

            <button class="green" type="submit">
                Submit Payment
            </button>

        </form>

    </div>
    """

    return page(
        "Payment",
        body,
        student=True,
        exam=exam,
        upi_id=upi_id
    )

# ============================================================
# PAYMENT QR
# ============================================================

@app.route("/student/payment/<int:exam_id>/qr")
def student_payment_qr(exam_id):

    guard = student_required()
    if guard:
        return guard

    if qrcode is None:

        flash("QR package is not installed.")

        return redirect(
            url_for(
                "student_payment",
                exam_id=exam_id
            )
        )

    exam = query_one(
        "SELECT * FROM exams WHERE id=?",
        (exam_id,)
    )

    if not exam:

        flash("Exam not found.")

        return redirect(
            url_for("student_dashboard")
        )

    # WTW Academy UPI ID
    upi_id = "kavitha94459@oksbi"

    amount = safe_float(
        exam["fee"]
    )

    # UPI payment URI
    uri = (
        "upi://pay?"
        "pa=" + upi_id +
        "&pn=WTW%20Academy"
        "&am=" + format(amount, ".2f") +
        "&cu=INR"
    )

    img = qrcode.make(uri)

    path = os.path.join(
        UPLOAD_DIR,
        "payment_qr_" +
        str(exam_id) +
        ".png"
    )

    img.save(path)

    return send_file(
        path,
        mimetype="image/png"
    )


# ============================================================
# STUDENT EXAM
# ============================================================

def load_attempt(attempt_id, student_id):

    return query_one("""
        SELECT a.*, e.exam_name, e.name AS old_exam_name,
               e.duration_minutes
        FROM exam_attempts a
        JOIN exams e
        ON e.id=a.exam_id
        WHERE a.id=?
        AND a.student_id=?
    """, (
        attempt_id,
        student_id
    ))


def attempt_questions(attempt):

    ids = [
        safe_int(x)
        for x in str(
            attempt["question_order"] or ""
        ).split(",")
        if str(x).strip()
    ]

    if not ids:
        return []

    placeholders = ",".join(
        "?" for _ in ids
    )

    rows = query_all(
        "SELECT * FROM questions "
        "WHERE id IN (" +
        placeholders +
        ")",
        ids
    )

    by_id = {
        row["id"]: row
        for row in rows
    }

    return [
        by_id[x]
        for x in ids
        if x in by_id
    ]


@app.route("/student/take/<int:attempt_id>")
def student_take(attempt_id):

    guard = student_required()
    if guard:
        return guard

    sid = session["student_db_id"]

    attempt = load_attempt(
        attempt_id,
        sid
    )

    if not attempt:
        flash("Exam attempt not found.")
        return redirect(url_for("student_dashboard"))

    if attempt["status"] == "Submitted":
        return redirect(
            url_for(
                "student_result_detail",
                attempt_id=attempt_id
            )
        )

    questions = attempt_questions(attempt)

    if not questions:
        flash("Questions not found.")
        return redirect(url_for("student_dashboard"))

    qindex = safe_int(
        request.args.get(
            "q",
            attempt["current_index"]
        ),
        0
    )

    qindex = max(
        0,
        min(qindex, len(questions) - 1)
    )

    question = questions[qindex]

    answer = query_one("""
        SELECT *
        FROM exam_answers
        WHERE attempt_id=?
        AND question_id=?
    """, (
        attempt_id,
        question["id"]
    ))

    selected = answer["selected_answer"] if answer else ""
    review = (
        bool(answer["marked_for_review"])
        if answer
        else False
    )

    elapsed = safe_int(
        attempt["elapsed_seconds"],
        0
    )

    if attempt["started_at"]:
        try:
            started = datetime.strptime(
                attempt["started_at"],
                "%Y-%m-%d %H:%M:%S"
            )

            elapsed = max(
                elapsed,
                int(
                    (
                        datetime.now() - started
                    ).total_seconds()
                )
            )
        except Exception:
            pass

    duration_seconds = (
        safe_int(attempt["duration_minutes"], 30)
        * 60
    )

    remaining = max(
        0,
        duration_seconds - elapsed
    )

    options = [
        ("A", question["option_a"]),
        ("B", question["option_b"]),
        ("C", question["option_c"]),
        ("D", question["option_d"]),
    ]

    body = """
    <div class="card">

        <button type="button"
                class="btn gray"
                onclick="toggleNav()">
            Hide / Show Navigation
        </button>

        <span class="badge orange"
              id="timer">
            Loading...
        </span>

    </div>

    <div class="exam-layout">

        <div class="card">

            <div class="muted">
                Question {{ qindex + 1 }}
                of {{ questions|length }}
            </div>

            <div class="question">
                {{ question.question_text }}
            </div>

            <form method="post"
                  action="{{ url_for(
                      'student_save_answer',
                      attempt_id=attempt.id
                  ) }}">

                <input type="hidden"
                       name="question_id"
                       value="{{ question.id }}">

                {% for letter,text in options %}

                <label class="option">
                    <input
                    type="{% if question.question_type == 'Multiple Correct' %}checkbox{% else %}radio{% endif %}"
                    name="answer"
                    value="{{ letter }}"
                    {% if letter in selected.split('|') %}
                    checked
                    {% endif %}>
                    <b>{{ letter }}.</b>
                    {{ text }}
                </label>

                {% endfor %}

                <label style="font-weight:400">
                    <input type="checkbox"
                           name="review"
                           value="1"
                           style="width:auto"
                           {% if review %}checked{% endif %}>
                    Mark for review
                </label>

                <button class="btn green">
                    Save & Next
                </button>

            </form>

            <div style="margin-top:12px">

                {% if qindex > 0 %}
                <a class="btn gray"
                   href="{{ url_for(
                       'student_take',
                       attempt_id=attempt.id
                   ) }}?q={{ qindex - 1 }}">
                   Previous
                </a>
                {% endif %}

                {% if qindex + 1 < questions|length %}
                <a class="btn"
                   href="{{ url_for(
                       'student_take',
                       attempt_id=attempt.id
                   ) }}?q={{ qindex + 1 }}">
                   Next
                </a>
                {% endif %}

                <form method="post"
                      action="{{ url_for(
                          'student_submit_exam',
                          attempt_id=attempt.id
                      ) }}"
                      style="display:inline"
                      onsubmit="return confirm('Submit exam?')">
                    <button class="btn red">
                        Submit Exam
                    </button>
                </form>

            </div>

            <p id="saveStatus"
               class="muted">
                Answers are saved automatically.
            </p>

        </div>

        <div class="card" id="navigationPanel">

            <h3>Question Navigator</h3>

            <div class="navigator">

            {% for q in questions %}

                {% set ans = answer_map.get(q.id) %}

                <a
                href="{{ url_for(
                    'student_take',
                    attempt_id=attempt.id
                ) }}?q={{ loop.index0 }}"
                class="
                {% if loop.index0 == qindex %}
                current
                {% endif %}
                {% if ans and ans.selected_answer %}
                answered
                {% endif %}
                {% if ans and ans.marked_for_review %}
                review
                {% endif %}
                ">
                    {{ loop.index }}
                </a>

            {% endfor %}

            </div>

        </div>

    </div>

    <script>
    let remaining={{ remaining }};

    function updateTimer(){

        let sec=remaining;

        let h=Math.floor(sec/3600);
        let m=Math.floor((sec%3600)/60);
        let s=sec%60;

        document.getElementById('timer').innerText =
            String(h).padStart(2,'0') + ':' +
            String(m).padStart(2,'0') + ':' +
            String(s).padStart(2,'0');

        if(remaining<=0){
            document.getElementById('timer').innerText='TIME UP';

            fetch(
                "{{ url_for(
                    'student_submit_exam',
                    attempt_id=attempt.id
                ) }}",
                {
                    method:"POST"
                }
            ).then(function(){
                window.location.href =
                    "{{ url_for(
                        'student_result_detail',
                        attempt_id=attempt.id
                    ) }}";
            });

            return;
        }

        remaining--;
    }

    setInterval(updateTimer,1000);
    updateTimer();

    function toggleNav(){

        const panel =
            document.getElementById('navigationPanel');

        if(panel.style.display==='none'){
            panel.style.display='block';
        }else{
            panel.style.display='none';
        }
    }

    const answerInputs =
        document.querySelectorAll(
            'input[name="answer"]'
        );

    answerInputs.forEach(function(input){

        input.addEventListener('change',function(){

            const data=new FormData();

            const checked =
                document.querySelectorAll(
                    'input[name="answer"]:checked'
                );

            const values=[];

            checked.forEach(function(x){
                values.push(x.value);
            });

            data.append(
                'question_id',
                '{{ question.id }}'
            );

            data.append(
                'answer',
                values.join('|')
            );

            const review =
                document.querySelector(
                    'input[name="review"]'
                );

            if(review && review.checked){
                data.append('review','1');
            }

            fetch(
                "{{ url_for(
                    'student_autosave',
                    attempt_id=attempt.id
                ) }}",
                {
                    method:'POST',
                    body:data
                }
            ).then(function(){
                document.getElementById(
                    'saveStatus'
                ).innerText='Saved ✓';
            });
        });
    });
    </script>
    """

    answer_rows = query_all("""
        SELECT *
        FROM exam_answers
        WHERE attempt_id=?
    """, (attempt_id,))

    answer_map = {
        row["question_id"]: row
        for row in answer_rows
    }

    return page(
        "Take Exam",
        body,
        student=True,
        attempt=attempt,
        questions=questions,
        question=question,
        qindex=qindex,
        options=options,
        selected=selected,
        review=review,
        remaining=remaining,
        answer_map=answer_map
    )
def answer_text(q, answer):

    if not answer:
        return "Not Answered"

    value = str(answer).strip()

    options = {
        "A": q["option_a"],
        "B": q["option_b"],
        "C": q["option_c"],
        "D": q["option_d"],
    }

    # Multiple Correct: A|C போன்ற answer
    if "|" in value:
        parts = [
            p.strip().upper()
            for p in value.split("|")
            if p.strip()
        ]

        return " | ".join(
            str(options.get(p, p))
            for p in parts
        )

    return str(options.get(value.upper(), value))

    return page(
        "Exam",
        body,
        student=True,
        attempt=attempt,
        questions=questions,
        question=question,
        qindex=qindex,
        selected=selected,
        review=review,
        options=options,
        remaining=remaining,
        answer_map=answer_map
    )


@app.route("/student/autosave/<int:attempt_id>",
           methods=["POST"])
def student_autosave(attempt_id):

    guard = student_required()
    if guard:
        return jsonify({"ok": False}), 401

    sid = session["student_db_id"]

    attempt = load_attempt(
        attempt_id,
        sid
    )

    if not attempt or attempt["status"] != "In Progress":
        return jsonify({"ok": False}), 400

    qid = safe_int(
        request.form.get("question_id")
    )

    answer = normalize_answer(
        request.form.get("answer", "")
    )

    review = (
        1
        if request.form.get("review") == "1"
        else 0
    )

    execute_write("""
        INSERT INTO exam_answers
        (attempt_id,question_id,selected_answer,
         marked_for_review,answered_at)
        VALUES (?,?,?,?,?)
        ON CONFLICT(attempt_id,question_id)
        DO UPDATE SET
            selected_answer=excluded.selected_answer,
            marked_for_review=excluded.marked_for_review,
            answered_at=excluded.answered_at
    """, (
        attempt_id,
        qid,
        answer,
        review,
        now_text()
    ))

    return jsonify({
        "ok": True
    })


@app.route("/student/exam/<int:attempt_id>/answer",
           methods=["POST"])
def student_save_answer(attempt_id):

    guard = student_required()
    if guard:
        return guard

    sid = session["student_db_id"]

    attempt = load_attempt(
        attempt_id,
        sid
    )

    if not attempt or attempt["status"] != "In Progress":
        return redirect(
            url_for("student_dashboard")
        )

    qid = safe_int(
        request.form.get("question_id")
    )

    answers = request.form.getlist("answer")

    if not answers:
        answer = ""
    else:
        answer = "|".join(
            sorted(
                set(
                    normalize_answer(x)
                    for x in answers
                    if normalize_answer(x)
                )
            )
        )

    review = (
        1
        if request.form.get("review") == "1"
        else 0
    )

    execute_write("""
        INSERT INTO exam_answers
        (attempt_id,question_id,selected_answer,
         marked_for_review,answered_at)
        VALUES (?,?,?,?,?)
        ON CONFLICT(attempt_id,question_id)
        DO UPDATE SET
            selected_answer=excluded.selected_answer,
            marked_for_review=excluded.marked_for_review,
            answered_at=excluded.answered_at
    """, (
        attempt_id,
        qid,
        answer,
        review,
        now_text()
    ))

    questions = attempt_questions(attempt)

    try:
        current_index = next(
            i for i,q in enumerate(questions)
            if q["id"] == qid
        )
    except StopIteration:
        current_index = safe_int(
            attempt["current_index"],
            0
        )

    next_index = min(
        current_index + 1,
        max(0, len(questions) - 1)
    )

    execute_write("""
        UPDATE exam_attempts
        SET current_index=?,
            current_question=?,
            last_question=?,
            elapsed_seconds=?
        WHERE id=?
    """, (
        next_index,
        questions[next_index]["id"]
        if questions else 0,
        current_index,
        safe_int(attempt["elapsed_seconds"], 0),
        attempt_id
    ))

    return redirect(
        url_for(
            "student_take",
            attempt_id=attempt_id
        ) +
        "?q=" +
        str(next_index)
    )


# ============================================================
# SUBMIT EXAM
# ============================================================

@app.route("/student/exam/<int:attempt_id>/submit",
           methods=["POST"])
@app.route("/student/submit/<int:attempt_id>",
           methods=["POST"])
def student_submit_exam(attempt_id):

    guard = student_required()
    if guard:
        return guard

    sid = session["student_db_id"]

    attempt = load_attempt(
        attempt_id,
        sid
    )

    if not attempt:
        flash("Attempt not found.")
        return redirect(url_for("student_dashboard"))

    if attempt["status"] == "Submitted":
        return redirect(
            url_for(
                "student_result_detail",
                attempt_id=attempt_id
            )
        )

    questions = attempt_questions(attempt)

    answers = query_all("""
        SELECT *
        FROM exam_answers
        WHERE attempt_id=?
    """, (
        attempt_id,
    ))

    answer_map = {
        a["question_id"]: a
        for a in answers
    }

    score = 0
    total = 0

    for q in questions:

        marks = safe_float(
            q["marks"],
            safe_float(
                attempt["total_marks"],
                0
            )
        )

        negative = safe_float(
            q["negative_mark"],
            0
        )

        total += marks

        answer = answer_map.get(q["id"])

        if not answer:
            continue

        selected = answer["selected_answer"]

        if not selected:
            continue

        correct = answer_is_correct(
            selected,
            q["correct_answer"],
            q["question_type"]
        )

        if correct:
            awarded = marks
            is_correct = 1
        else:
            awarded = -negative
            is_correct = 0

        score += awarded

        execute_write("""
            UPDATE exam_answers
            SET is_correct=?,
                marks_awarded=?
            WHERE id=?
        """, (
            is_correct,
            awarded,
            answer["id"]
        ))

    elapsed = safe_int(
        attempt["elapsed_seconds"],
        0
    )

    if attempt["started_at"]:
        try:
            started = datetime.strptime(
                attempt["started_at"],
                "%Y-%m-%d %H:%M:%S"
            )

            elapsed = max(
                elapsed,
                int(
                    (
                        datetime.now() - started
                    ).total_seconds()
                )
            )
        except Exception:
            pass

    execute_write("""
        UPDATE exam_attempts
        SET status='Submitted',
            submitted_at=?,
            elapsed_seconds=?,
            score=?,
            total_marks=?
        WHERE id=?
    """, (
        now_text(),
        elapsed,
        score,
        total,
        attempt_id
    ))

    return redirect(
        url_for(
            "student_result_detail",
            attempt_id=attempt_id
        )
    )


# ============================================================
# STUDENT RESULTS
# ============================================================

@app.route("/student/results")
def student_results():

    guard = student_required()
    if guard:
        return guard

    sid = session["student_db_id"]

    results = query_all("""
        SELECT a.*,
               e.exam_name,
               e.name AS old_exam_name
        FROM exam_attempts a
        JOIN exams e
        ON e.id = a.exam_id
        WHERE a.student_id = ?
        AND a.status = 'Submitted'
        ORDER BY a.id DESC
    """, (sid,))

    body = """
    <div class="card">

        <h1>My Results</h1>

        <div class="table-scroll">

        <table>

        <tr>
            <th>Date</th>
            <th>Exam</th>
            <th>Attempt</th>
            <th>Score</th>
            <th>Total</th>
            <th>Action</th>
        </tr>

        {% for r in results %}

        <tr>

            <td>
                {{ r.submitted_at }}
            </td>

            <td>
                {{ r.exam_name or r.old_exam_name }}
            </td>

            <td>
                {{ r.attempt_no }}
            </td>

            <td>
                {{ r.score }}
            </td>

            <td>
                {{ r.total_marks }}
            </td>

            <td>
                <a class="btn small"
                   href="{{ url_for(
                       'student_result_detail',
                       attempt_id=r.id
                   ) }}">
                   Review
                </a>
            </td>

        </tr>

        {% endfor %}

        </table>

        </div>

    </div>
    """

    return page(
        "My Results",
        body,
        student=True,
        results=results
    )


@app.route("/student/results/<int:attempt_id>")
def student_result_detail(attempt_id):

    guard = student_required()
    if guard:
        return guard

    sid = session["student_db_id"]

    attempt = load_attempt(
        attempt_id,
        sid
    )

    if not attempt:
        flash("Result not found.")
        return redirect(
            url_for("student_results")
        )

    if attempt["status"] != "Submitted":
        return redirect(
            url_for(
                "student_take",
                attempt_id=attempt_id
            )
        )

    questions = attempt_questions(attempt)

    answers = query_all("""
        SELECT *
        FROM exam_answers
        WHERE attempt_id = ?
    """, (attempt_id,))

    answer_map = {
        a["question_id"]: a
        for a in answers
    }

    body = """
    <div class="card">

        <h1>
            {{ attempt.exam_name or attempt.old_exam_name }}
        </h1>

        <p>
            Score:
            <b>{{ attempt.score }}</b>
            /
            {{ attempt.total_marks }}
        </p>

        <p>
            Submitted:
            {{ attempt.submitted_at }}
        </p>

    </div>


    {% for q in questions %}

    {% set a = answer_map.get(q.id) %}


    <div class="card">

        <div class="question">

            <b>{{ loop.index }}.</b>

            {{ q.question_text }}

        </div>

<!-- STUDENT ANSWER -->

<p>
    <b>Your Answer:</b>

    {% if a and a.selected_answer %}

        {% for raw_ans in a.selected_answer.split("|") %}

            {% set ans = raw_ans|trim|upper|replace('"', '')|replace("'", "") %}

            {% if ans == "A" %}
                {{ q.option_a }}
            {% elif ans == "B" %}
                {{ q.option_b }}
            {% elif ans == "C" %}
                {{ q.option_c }}
            {% elif ans == "D" %}
                {{ q.option_d }}
            {% else %}
                {{ raw_ans }}
            {% endif %}

            {% if not loop.last %}, {% endif %}

        {% endfor %}

    {% else %}

        Not Answered

    {% endif %}
</p>


<!-- CORRECT ANSWER -->

<p>
    <b>Correct Answer:</b>

    {% if q.correct_answer %}

        {% for raw_ans in q.correct_answer.split("|") %}

            {% set ans = raw_ans|trim|upper|replace('"', '')|replace("'", "") %}

            {% if ans == "A" %}
                {{ q.option_a }}
            {% elif ans == "B" %}
                {{ q.option_b }}
            {% elif ans == "C" %}
                {{ q.option_c }}
            {% elif ans == "D" %}
                {{ q.option_d }}
            {% else %}
                {{ raw_ans }}
            {% endif %}

            {% if not loop.last %}, {% endif %}

        {% endfor %}

    {% else %}

        Not Available

    {% endif %}
</p>



        <!-- RESULT STATUS -->

        {% if a %}

            {% if a.is_correct %}

                <span class="badge green">
                    Correct
                </span>

            {% else %}

                <span class="badge red">
                    Incorrect
                </span>

            {% endif %}

        {% else %}

            <span class="badge orange">
                Not Answered
            </span>

        {% endif %}


        <!-- MARKS -->

        <p>

            <b>Marks Awarded:</b>

            {{ a.marks_awarded if a else 0 }}

        </p>


        <!-- EXPLANATION -->

        {% if q.explanation %}

            <p class="muted">

                <b>Explanation:</b>

                {{ q.explanation }}

            </p>

        {% endif %}


    </div>

    {% endfor %}


    <a class="btn"
       href="{{ url_for('student_results') }}">

        Back to Results

    </a>
    """

    return page(
        "Result Review",
        body,
        student=True,
        attempt=attempt,
        questions=questions,
        answer_map=answer_map
    )
# ============================================================
# ERROR HANDLERS
# ============================================================

@app.errorhandler(404)
def not_found(error):
    body = """
    <div class="card center">
        <h1>404</h1>
        <p>Page not found.</p>
        <a class="btn"
           href="{{ url_for('home') }}">
           Home
        </a>
    </div>
    """

    return page(
        "Page Not Found",
        body,
        nav=False
    ), 404


@app.errorhandler(500)
def server_error(error):
    body = """
    <div class="card center">
        <h1>500</h1>
        <p>Something went wrong.</p>
        <a class="btn"
           href="{{ url_for('home') }}">
           Home
        </a>
    </div>
    """

    return page(
        "Server Error",
        body,
        nav=False
    ), 500


# ============================================================
# INITIALIZE
# ============================================================

try:
    migrate_database()
    migrate_local_sqlite_to_postgres()
except Exception as init_error:
    print("Database initialization error:", init_error)
    raise


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":
    print("")
    print("==========================================")
    print("WTW Academy")
    print("==========================================")
    print("Database:", "PostgreSQL (DATABASE_URL)" if USE_POSTGRES else DB_PATH)
    print("URL: http://127.0.0.1:5000")
    print("==========================================")
    print("")

    app.run(
        host="127.0.0.1",
        port=5000,
        debug=False,
        use_reloader=False,
        threaded=True
    )



